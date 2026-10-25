import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook(r"e:\MY PROJECT\Freelance Jobs\خبير لينكد ان\B2B_Training_Sales_Riyadh.xlsx")
ws = wb.active

center_aligned_text = Alignment(horizontal="center", vertical="center")
header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
header_font = Font(bold=True)

for row in ws.iter_rows():
    for cell in row:
        cell.alignment = center_aligned_text
        if cell.row == 1:
            cell.font = header_font
            cell.fill = header_fill

for col in ws.columns:
    max_length = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        try:
            if cell.value and len(str(cell.value)) > max_length:
                max_length = len(str(cell.value))
        except:
            pass
    adjusted_width = max_length + 2
    if adjusted_width > 45:
        adjusted_width = 45
    ws.column_dimensions[col_letter].width = adjusted_width

wb.save(r"e:\MY PROJECT\Freelance Jobs\خبير لينكد ان\B2B_Training_Sales_Riyadh.xlsx")
print("Formatting complete")
