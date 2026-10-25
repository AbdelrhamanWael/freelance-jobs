import docx

doc = docx.Document(r"e:\MY PROJECT\Freelance Jobs\خبير لينكد ان\مرشحين-مبيعات-وتسويق.docx")
with open(r"e:\MY PROJECT\Freelance Jobs\خبير لينكد ان\template_structure.txt", "w", encoding="utf-8") as f:
    f.write("Paragraphs:\n")
    for para in doc.paragraphs:
        f.write(para.text + "\n")
    f.write("\nTables:\n")
    for table in doc.tables:
        for row in table.rows:
            f.write(str([cell.text for cell in row.cells]) + "\n")
        f.write("---\n")
