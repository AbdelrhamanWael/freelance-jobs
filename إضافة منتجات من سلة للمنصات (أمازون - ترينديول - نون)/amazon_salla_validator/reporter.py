"""
Multi-format report generator: Console table, CSV, JSON, and styled Excel.
"""

import json
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any
import pandas as pd

from .config import ProductComparison


import sys


def safe_print(text: str):
    """Prints text safely without throwing UnicodeEncodeError on legacy consoles."""
    try:
        print(text)
    except UnicodeEncodeError:
        encoding = sys.stdout.encoding or "ascii"
        fallback_text = text.encode(encoding, errors="replace").decode(encoding)
        print(fallback_text)


class ValidationReporter:
    """Generates console, CSV, JSON, and Excel reports."""

    def __init__(self, comparisons: List[ProductComparison], output_dir: Path):
        self.comparisons = comparisons
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        if hasattr(sys.stdout, "reconfigure"):
            try:
                sys.stdout.reconfigure(encoding="utf-8")
            except Exception:
                pass

    def get_summary_stats(self) -> Dict[str, Any]:
        """Calculates validation metrics."""
        total = len(self.comparisons)
        matched = sum(1 for c in self.comparisons if c.overall_match)
        mismatched = total - matched
        not_found = sum(1 for c in self.comparisons if c.amazon_title == "NOT FOUND")

        total_fields = sum(len(c.diffs) for c in self.comparisons)
        matched_fields = sum(
            sum(1 for d in c.diffs if d.matched)
            for c in self.comparisons
        )
        accuracy_pct = (matched_fields / total_fields * 100) if total_fields else 0.0

        return {
            "total_products": total,
            "overall_matched": matched,
            "overall_mismatched": mismatched,
            "not_found_on_amazon": not_found,
            "total_fields_checked": total_fields,
            "matched_fields": matched_fields,
            "accuracy_percentage": round(accuracy_pct, 2),
            "generated_at": datetime.now().isoformat(),
        }

    def print_console_summary(self):
        """Prints a human-readable console report."""
        stats = self.get_summary_stats()

        safe_print("\n" + "=" * 80)
        safe_print("  SALLA VS AMAZON PRODUCT DATA VALIDATION SUMMARY")
        safe_print("=" * 80)
        safe_print(f"  Total Products Checked:    {stats['total_products']}")
        safe_print(f"  Fully Matched:             {stats['overall_matched']} (✓)")
        safe_print(f"  Flagged with Mismatches:   {stats['overall_mismatched']} (✗)")
        safe_print(f"  Not Found in Inventory:    {stats['not_found_on_amazon']}")
        safe_print(f"  Field-Level Accuracy:      {stats['accuracy_percentage']}%")
        safe_print("=" * 80 + "\n")

        # Print detailed discrepancy log
        for comp in self.comparisons:
            status_symbol = "✓ PASS" if comp.overall_match else "✗ REVIEW NEEDED"
            safe_print(f"[{status_symbol}] SKU: {comp.sku} | ASIN: {comp.asin} | Name: {comp.salla_name[:40]}")
            for diff in comp.diffs:
                icon = diff.status_icon
                note = f" -> {diff.difference_note}" if diff.difference_note else ""
                safe_print(f"    {icon} {diff.field_name:<16}: Expected '{diff.expected_val}' | Actual '{diff.actual_val}'{note}")
            safe_print("-" * 80)

    def to_flattened_rows(self) -> List[Dict[str, Any]]:
        """Flattens comparisons into individual field rows for tabular export."""
        rows = []
        for comp in self.comparisons:
            for diff in comp.diffs:
                rows.append({
                    "Product Identifier (SKU)": comp.sku,
                    "Amazon ASIN": comp.asin,
                    "Product Name (Salla)": comp.salla_name,
                    "Amazon Title": comp.amazon_title,
                    "Field Name": diff.field_name,
                    "Expected Value (Salla)": diff.expected_val,
                    "Actual Value (Amazon)": diff.actual_val,
                    "Match Status": diff.status_icon,
                    "Severity": diff.severity,
                    "Difference Notes": diff.difference_note,
                })
        return rows

    def export_csv(self) -> Path:
        """Exports validation report to CSV."""
        rows = self.to_flattened_rows()
        df = pd.DataFrame(rows)
        file_path = self.output_dir / f"amazon_validation_report_{self.timestamp}.csv"
        df.to_csv(file_path, index=False, encoding="utf-8-sig")

        # Also save latest alias safely
        latest_path = self.output_dir / "amazon_validation_report_latest.csv"
        try:
            df.to_csv(latest_path, index=False, encoding="utf-8-sig")
        except (PermissionError, IOError):
            pass

        return file_path

    def export_json(self) -> Path:
        """Exports validation report to JSON."""
        stats = self.get_summary_stats()
        data = {
            "summary": stats,
            "products": [
                {
                    "sku": c.sku,
                    "asin": c.asin,
                    "salla_name": c.salla_name,
                    "amazon_title": c.amazon_title,
                    "overall_match": c.overall_match,
                    "error_message": c.error_message,
                    "diffs": [
                        {
                            "field_name": d.field_name,
                            "expected_val": d.expected_val,
                            "actual_val": d.actual_val,
                            "matched": d.matched,
                            "status_icon": d.status_icon,
                            "severity": d.severity,
                            "difference_note": d.difference_note,
                        }
                        for d in c.diffs
                    ]
                }
                for c in self.comparisons
            ]
        }

        file_path = self.output_dir / f"amazon_validation_report_{self.timestamp}.json"
        with open(file_path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

        latest_path = self.output_dir / "amazon_validation_report_latest.json"
        try:
            with open(latest_path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except (PermissionError, IOError):
            pass

        return file_path

    def export_excel(self) -> Path:
        """Exports styled Excel workbook with conditional coloring."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        rows = self.to_flattened_rows()
        file_path = self.output_dir / f"amazon_validation_report_{self.timestamp}.xlsx"
        latest_path = self.output_dir / "amazon_validation_report_latest.xlsx"

        wb = openpyxl.Workbook()

        # Sheet 1: Detailed Discrepancies
        ws_details = wb.active
        ws_details.title = "Field Discrepancies"

        headers = [
            "Product Identifier (SKU)", "Amazon ASIN", "Product Name (Salla)",
            "Field Name", "Expected Value (Salla)", "Actual Value (Amazon)",
            "Match Status", "Severity", "Difference Notes"
        ]
        ws_details.append(headers)

        # Header styling
        header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="D1D5DB"),
            right=Side(style="thin", color="D1D5DB"),
            top=Side(style="thin", color="D1D5DB"),
            bottom=Side(style="thin", color="D1D5DB")
        )

        for col_idx, _ in enumerate(headers, 1):
            cell = ws_details.cell(1, col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Color fills for status
        green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        yellow_fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")

        # Populate rows
        for row in rows:
            row_vals = [
                row["Product Identifier (SKU)"],
                row["Amazon ASIN"],
                row["Product Name (Salla)"],
                row["Field Name"],
                row["Expected Value (Salla)"],
                row["Actual Value (Amazon)"],
                row["Match Status"],
                row["Severity"],
                row["Difference Notes"],
            ]
            ws_details.append(row_vals)
            current_row = ws_details.max_row

            # Apply cell formatting
            status_val = row["Match Status"]
            severity_val = row["Severity"]

            if status_val == "✓" or severity_val == "MATCH":
                row_fill = green_fill
            elif status_val == "✗" or severity_val == "MISMATCH":
                row_fill = red_fill
            else:
                row_fill = yellow_fill

            for col_idx in range(1, len(headers) + 1):
                c = ws_details.cell(current_row, col_idx)
                c.border = thin_border
                c.font = Font(name="Segoe UI", size=10)
                if col_idx in [7, 8]:  # Status & Severity column
                    c.fill = row_fill
                    c.alignment = Alignment(horizontal="center", vertical="center")

        # Auto-adjust column widths
        for col in ws_details.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws_details.column_dimensions[col_letter].width = max(12, min(max_len + 3, 50))

        # Sheet 2: Executive Summary
        ws_summary = wb.create_sheet(title="Executive Summary")
        stats = self.get_summary_stats()

        ws_summary.append(["Salla vs Amazon Catalog Validation Summary"])
        ws_summary.append([])
        ws_summary.append(["Metric", "Value"])

        summary_rows = [
            ("Total Products Audited", stats["total_products"]),
            ("Products Fully Matched (✓)", stats["overall_matched"]),
            ("Products with Discrepancies (✗)", stats["overall_mismatched"]),
            ("Products Not Found in Amazon", stats["not_found_on_amazon"]),
            ("Total Fields Evaluated", stats["total_fields_checked"]),
            ("Matched Fields Count", stats["matched_fields"]),
            ("Catalog Accuracy Rate (%)", f"{stats['accuracy_percentage']}%"),
            ("Report Timestamp", stats["generated_at"]),
        ]

        for k, v in summary_rows:
            ws_summary.append([k, v])

        # Style summary header
        ws_summary.cell(1, 1).font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
        ws_summary.cell(3, 1).font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        ws_summary.cell(3, 1).fill = header_fill
        ws_summary.cell(3, 2).font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        ws_summary.cell(3, 2).fill = header_fill

        ws_summary.column_dimensions["A"].width = 35
        ws_summary.column_dimensions["B"].width = 25

        wb.save(file_path)
        try:
            wb.save(latest_path)
        except (PermissionError, IOError):
            pass
        return file_path

    def export_all(self) -> Dict[str, Path]:
        """Generates console output and all export files."""
        self.print_console_summary()
        csv_path = self.export_csv()
        json_path = self.export_json()
        excel_path = self.export_excel()

        return {
            "csv": csv_path,
            "json": json_path,
            "excel": excel_path,
        }
