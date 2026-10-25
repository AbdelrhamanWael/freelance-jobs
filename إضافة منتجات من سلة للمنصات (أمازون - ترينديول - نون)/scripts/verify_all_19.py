import openpyxl
import sys

if sys.stdout.encoding != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except:
        pass

wb = openpyxl.load_workbook('data/exports/amazon_drafts_completed_ready_to_upload.xlsm', data_only=True)
ws = wb['Template']

cols_to_check = {
    10: "Product Id Type",
    11: "Product Id",
    23: "Model Number",
    68: "Number of Items",
    96: "Ingredients",
    110: "Is Product Expirable",
    111: "Product Expiration Type",
    115: "Unit Count",
    116: "Unit Count Type",
    149: "Temperature Rating",
    521: "Item Condition",
    523: "List Price Currency",
    524: "List Price with Tax",
    548: "Fulfillment Channel Code (SA)",
    549: "Quantity (SA)",
    550: "Handling Time (SA)",
    553: "Your Price SAR",
    663: "Dangerous Goods Regulations",
    691: "Contains Battery or Cell"
}

missing_per_col = {col: 0 for col in cols_to_check}

for r in range(7, ws.max_row + 1):
    for col in cols_to_check:
        val = ws.cell(r, col).value
        if val is None or str(val).strip() == '':
            missing_per_col[col] += 1

print(f"Verified {ws.max_row - 6} products:")
all_ok = True
for col, name in cols_to_check.items():
    cnt = missing_per_col[col]
    sample_val = ws.cell(7, col).value
    print(f"  Col {col} ({name}): missing = {cnt} | sample = '{sample_val}'")
    if cnt > 0:
        all_ok = False

if all_ok:
    print("\n🎉 ALL 19 ATTRIBUTES ARE 100% COMPLETE FOR ALL 270 PRODUCTS!")
else:
    print("\n⚠️ Some columns still have missing values!")
