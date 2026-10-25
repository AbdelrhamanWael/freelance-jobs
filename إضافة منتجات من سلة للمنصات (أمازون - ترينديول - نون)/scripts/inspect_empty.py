import openpyxl

wb = openpyxl.load_workbook('data/amazon_upload_ready_290_products_2026-09-13_drafts.xlsm', data_only=True)
ws = wb['Template']

empty_cols_row7 = []
for c in range(1, ws.max_column + 1):
    val = ws.cell(7, c).value
    lbl = ws.cell(4, c).value
    key = ws.cell(5, c).value
    if val is None and lbl is not None:
        empty_cols_row7.append((c, lbl, key))

print(f"Total empty columns in Row 7: {len(empty_cols_row7)}")
# Print the ones that might be required or common
common = ['product id', 'model', 'brand', 'price', 'quantity', 'image', 'description', 'temperature', 'expiration', 'ingredients', 'items', 'unit', 'country']
for c, lbl, key in empty_cols_row7:
    lbl_s = str(lbl).lower()
    if any(k in lbl_s for k in common):
        print(f"  Col {c}: {lbl} (key: {key})")
