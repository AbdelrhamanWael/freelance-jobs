import openpyxl
import pandas as pd

# Load Salla cleaned database
df_clean = pd.read_pickle('data/clean/amazon_ready.pkl')
clean_lookup = {str(row['sku']).strip(): row for _, row in df_clean.iterrows()}
print(f"Loaded {len(clean_lookup)} SKUs from amazon_ready.pkl")

# Load template workbook
wb = openpyxl.load_workbook('data/amazon_upload_ready_290_products_2026-09-13_drafts.xlsm')
ws = wb['Template']

print(f"Template total rows: {ws.max_row}")

# Check matching SKUs
matched = 0
unmatched = []
for r in range(7, ws.max_row + 1):
    sku = str(ws.cell(r, 2).value or '').strip()
    if sku in clean_lookup:
        matched += 1
    else:
        unmatched.append((r, sku))

print(f"Matched SKUs in template: {matched} / {ws.max_row - 6}")
if unmatched:
    print(f"Unmatched SKUs ({len(unmatched)}): {unmatched[:5]}")
