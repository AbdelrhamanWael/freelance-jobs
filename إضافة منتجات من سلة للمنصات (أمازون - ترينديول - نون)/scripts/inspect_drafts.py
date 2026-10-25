import openpyxl

wb = openpyxl.load_workbook('data/amazon_upload_ready_290_products_2026-09-13_drafts.xlsm', data_only=True)
ws = wb['Template']
row4 = [ws.cell(4, c).value for c in range(1, ws.max_column + 1)]
row5 = [ws.cell(5, c).value for c in range(1, ws.max_column + 1)]

targets = ['temperature', 'number of items', 'ingredients', 'model', 'highlight', 'expiration', 'unit count', 'allergen', 'sku', 'brand', 'product id']
found = {}
for idx, (h4, h5) in enumerate(zip(row4, row5), 1):
    h4_str = str(h4).lower() if h4 else ''
    h5_str = str(h5).lower() if h5 else ''
    for t in targets:
        if t in h4_str or t in h5_str:
            found.setdefault(t, []).append((idx, h4, h5))

for k, v in found.items():
    print(f"=== Target: {k} ({len(v)} matches) ===")
    for col_idx, h4, h5 in v[:4]:
        print(f"  Col {col_idx}: Label='{h4}' | Key='{h5}'")
