import openpyxl

wb = openpyxl.load_workbook('data/amazon_upload_ready_290_products_2026-09-13_drafts.xlsm', data_only=True)
ws_vals = wb['Valid Values']
print("Valid Values max row:", ws_vals.max_row, "max col:", ws_vals.max_column)

# Row 1 or 2 contains attribute names
headers = [ws_vals.cell(1, c).value or ws_vals.cell(2, c).value for c in range(1, ws_vals.max_column + 1)]

for col_idx, h in enumerate(headers, 1):
    h_str = str(h).lower()
    if any(k in h_str for k in ['temperature', 'expiration', 'unit count type', 'allergen']):
        vals = [ws_vals.cell(r, col_idx).value for r in range(3, min(20, ws_vals.max_row + 1))]
        vals = [v for v in vals if v is not None]
        print(f"Col {col_idx} '{h}': {vals[:5]}")
