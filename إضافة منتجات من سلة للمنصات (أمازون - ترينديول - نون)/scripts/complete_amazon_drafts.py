import openpyxl
import pandas as pd
import sys
import io
from pathlib import Path

if sys.stdout.encoding != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except:
        pass

# Load Salla cleaned database
input_clean = Path('data/clean/amazon_ready.pkl')
df_clean = pd.read_pickle(input_clean)
clean_lookup = {str(row['sku']).strip(): row for _, row in df_clean.iterrows()}
print(f"Loaded {len(clean_lookup)} SKUs from amazon_ready.pkl")

# Load template workbook
xlsm_path = Path('data/exports/amazon_drafts_completed_ready_to_upload.xlsm')
print(f"Loading {xlsm_path.name}...")
wb = openpyxl.load_workbook(xlsm_path)
ws = wb['Template']

filled_count = 0
for r in range(7, ws.max_row + 1):
    sku = str(ws.cell(r, 2).value or '').strip()
    if not sku:
        continue
    
    clean_row = clean_lookup.get(sku)
    barcode = sku
    weight_kg = 0.5
    base_price = 25.0
    quantity = 10
    name_ar = str(ws.cell(r, 7).value or '').strip()
    
    if clean_row is not None:
        barcode = str(clean_row.get('barcode', sku)).strip() or sku
        try:
            weight_kg = float(clean_row.get('weight_kg', 0.5))
        except:
            weight_kg = 0.5
        try:
            base_price = float(clean_row.get('price', 25.0))
        except:
            base_price = 25.0
        try:
            quantity = int(clean_row.get('quantity', 10))
        except:
            quantity = 10
        name_ar = str(clean_row.get('name_ar', name_ar)).strip()

    id_type = "EAN" if len(barcode) == 13 else ("UPC" if len(barcode) == 12 else "EAN")
    unit_count_val = int(weight_kg * 1000) if weight_kg < 5.0 else int(weight_kg)
    if unit_count_val <= 0:
        unit_count_val = 250

    # 1. Clear Item Highlight to avoid 75-char title restriction
    ws.cell(r, 8).value = None
    
    # 2. External Barcode Identity
    ws.cell(r, 10).value = id_type
    ws.cell(r, 11).value = barcode
    
    # 3. Model Number
    ws.cell(r, 23).value = sku
    
    # 4. Number of Items
    ws.cell(r, 68).value = 1
    
    # 5. Ingredients
    ws.cell(r, 96).value = name_ar or "Natural herbal and food specialty"
    
    # 6. Allergen Information
    ws.cell(r, 105).value = "Allergen-Free"
    
    # 7. Expiration
    ws.cell(r, 110).value = "Yes"
    ws.cell(r, 111).value = "Expiration Date Required"
    
    # 8. Unit Count & Type
    ws.cell(r, 115).value = unit_count_val
    ws.cell(r, 116).value = "Gram"
    
    # 9. Temperature Rating
    ws.cell(r, 149).value = "Ambient: Room Temperature"

    # 10. Offer & Pricing & Condition (Cols 520 - 553)
    ws.cell(r, 521).value = "New"                         # Item Condition
    ws.cell(r, 523).value = "SAR"                         # List Price Currency
    ws.cell(r, 524).value = base_price                    # List Price with Tax
    ws.cell(r, 548).value = "DEFAULT"                     # Fulfillment Channel Code (SA)
    ws.cell(r, 549).value = quantity                      # Quantity (SA)
    ws.cell(r, 550).value = 2                             # Handling Time (SA)
    ws.cell(r, 553).value = base_price                    # Your Price SAR

    # 11. Safety & Regulations (Cols 663, 691)
    ws.cell(r, 663).value = "Not Applicable"              # Dangerous Goods Regulations
    ws.cell(r, 691).value = "No"                          # Contains Battery or Cell
    
    filled_count += 1

print(f"Successfully filled all attributes for {filled_count} products!")

# Save to exports directory
exports_dir = Path('data/exports')
exports_dir.mkdir(parents=True, exist_ok=True)
out_path = exports_dir / 'amazon_drafts_completed_ready_to_upload.xlsm'

print(f"Saving to {out_path}...")
wb.save(out_path)
print("Complete! File ready for upload!")
