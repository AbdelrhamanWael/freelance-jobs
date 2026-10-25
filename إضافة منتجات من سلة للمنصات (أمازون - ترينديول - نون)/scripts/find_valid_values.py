import openpyxl

wb = openpyxl.load_workbook('data/amazon_upload_ready_290_products_2026-09-13_drafts.xlsm', data_only=True)
ws_vals = wb['Valid Values']

for col_idx in range(1, ws_vals.max_column + 1):
    h1 = str(ws_vals.cell(1, col_idx).value or '').strip()
    h2 = str(ws_vals.cell(2, col_idx).value or '').strip()
    full_h = f"{h1} {h2}".lower()
    
    if any(k in full_h for k in ['temperature rating', 'product expiration type', 'is product expirable', 'unit count type']):
        vals = [ws_vals.cell(r, col_idx).value for r in range(3, min(25, ws_vals.max_row + 1))]
        vals = [v for v in vals if v is not None]
        print(f"Col {col_idx}: '{h1}' / '{h2}': {vals}")
