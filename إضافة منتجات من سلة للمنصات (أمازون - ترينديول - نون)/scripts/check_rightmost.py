import openpyxl

wb = openpyxl.load_workbook('data/exports/amazon_drafts_completed_ready_to_upload.xlsm', data_only=True)
ws = wb['Template']

for c in range(650, ws.max_column + 1):
    lbl = ws.cell(4, c).value
    if lbl:
        print(f"Col {c}: {lbl}")
