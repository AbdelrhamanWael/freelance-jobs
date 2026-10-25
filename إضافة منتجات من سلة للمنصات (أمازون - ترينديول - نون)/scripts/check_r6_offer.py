import openpyxl

wb = openpyxl.load_workbook('data/exports/amazon_drafts_completed_ready_to_upload.xlsm', data_only=True)
ws = wb['Template']
for c in [520, 521, 522, 523, 524, 526, 548, 549, 550, 553]:
    lbl = ws.cell(4, c).value
    v6 = ws.cell(6, c).value
    print(f"Col {c} ({lbl}): Row 6 value = '{v6}'")
