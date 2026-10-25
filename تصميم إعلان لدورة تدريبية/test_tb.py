import docx
from docx.shared import Inches
from docx.oxml import parse_xml
import win32com.client as win32
import os
import pymupdf

doc = docx.Document()
section = doc.sections[0]

page_w = Inches(11.7533)
page_h = Inches(9.18)
section.page_width = page_w
section.page_height = page_h
section.top_margin = Inches(0)
section.bottom_margin = Inches(0)
section.left_margin = Inches(0)
section.right_margin = Inches(0)

# Add background image in header
header = section.header
p_hdr = header.paragraphs[0]
run_hdr = p_hdr.add_run()
pic = run_hdr.add_picture('exact_bg_template.png', width=page_w, height=page_h)

inline = run_hdr._r.xpath('.//wp:inline')[0]
extent = inline.xpath('.//wp:extent')[0]
cx = extent.get('cx')
cy = extent.get('cy')
graphic = inline.xpath('.//a:graphic')[0]

anchor_xml = f'''
<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
           xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
           distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
    <wp:simplePos x="0" y="0"/>
    <wp:positionH relativeFrom="page">
        <wp:posOffset>0</wp:posOffset>
    </wp:positionH>
    <wp:positionV relativeFrom="page">
        <wp:posOffset>0</wp:posOffset>
    </wp:positionV>
    <wp:extent cx="{cx}" cy="{cy}"/>
    <wp:effectExtent l="0" t="0" r="0" b="0"/>
    <wp:wrapNone/>
    <wp:docPr id="1" name="Background"/>
    <wp:cNvGraphicFramePr/>
</wp:anchor>
'''
anchor_el = parse_xml(anchor_xml)
anchor_el.append(graphic)
inline.getparent().replace(inline, anchor_el)

# Helper to convert pixels (at 1763x1377 over 11.7533 x 9.18 in) to EMUs
# 1 inch = 914400 EMUs
# 1763 px = 11.7533 inches -> 1 px = 914400 * 11.7533 / 1763 = 6096 EMUs
PX_TO_EMU = int(914400 * 11.7533 / 1763)

def add_floating_textbox(doc, x_px, y_px, w_px, h_px, text, font_size_pt, color_hex, bold=True, align="center", box_id=10):
    x_emu = int(x_px * PX_TO_EMU)
    y_emu = int(y_px * PX_TO_EMU)
    w_emu = int(w_px * PX_TO_EMU)
    h_emu = int(h_px * PX_TO_EMU)
    sz_half = int(font_size_pt * 2)
    b_tag = "<w:b/>" if bold else ""
    
    tb_xml = f'''
    <w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
         xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
         xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
         xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
      <w:pPr>
        <w:spacing w:before="0" w:after="0" w:line="0" w:lineRule="auto"/>
      </w:pPr>
      <w:r>
        <w:drawing>
          <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251658240" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
            <wp:simplePos x="0" y="0"/>
            <wp:positionH relativeFrom="page">
              <wp:posOffset>{x_emu}</wp:posOffset>
            </wp:positionH>
            <wp:positionV relativeFrom="page">
              <wp:posOffset>{y_emu}</wp:posOffset>
            </wp:positionV>
            <wp:extent cx="{w_emu}" cy="{h_emu}"/>
            <wp:effectExtent l="0" t="0" r="0" b="0"/>
            <wp:wrapNone/>
            <wp:docPr id="{box_id}" name="TextBox_{box_id}"/>
            <wp:cNvGraphicFramePr/>
            <a:graphic>
              <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
                <wps:wsp>
                  <wps:spPr>
                    <a:xfrm>
                      <a:off x="0" y="0"/>
                      <a:ext cx="{w_emu}" cy="{h_emu}"/>
                    </a:xfrm>
                    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
                    <a:noFill/>
                    <a:ln><a:noFill/></a:ln>
                  </wps:spPr>
                  <wps:txbx>
                    <w:txbxContent>
                      <w:p>
                        <w:pPr>
                          <w:bidi/>
                          <w:jc w:val="{align}"/>
                          <w:spacing w:before="0" w:after="0"/>
                        </w:pPr>
                        <w:r>
                          <w:rPr>
                            <w:rtl/>
                            <w:rFonts w:ascii="Segoe UI" w:hAnsi="Segoe UI" w:cs="Segoe UI"/>
                            <w:b/>
                            <w:bCs/>
                            <w:color w:val="{color_hex}"/>
                            <w:sz w:val="{sz_half}"/>
                            <w:szCs w:val="{sz_half}"/>
                          </w:rPr>
                          <w:t>{text}</w:t>
                        </w:r>
                      </w:p>
                    </w:txbxContent>
                  </wps:txbx>
                  <wps:bodyPr lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"/>
                </wps:wsp>
              </a:graphicData>
            </a:graphic>
          </wp:anchor>
        </w:drawing>
      </w:r>
    </w:p>
    '''
    p_el = parse_xml(tb_xml)
    doc._body._body.append(p_el)

# Add Test Header Textboxes based on Playwright coordinates:
# Institute: x=565.5, y=168.3, w=632.0, h=39.0
add_floating_textbox(doc, 565.5, 160.0, 632.0, 50.0, "معهد مئة وثمانون درجة العالي للتدريب", 16, "3BA4A6", bold=True, align="center", box_id=101)
# Title: x=565.5, y=215.3, w=632.0, h=36.5
add_floating_textbox(doc, 565.5, 205.0, 632.0, 60.0, "دورة التأسيس لاختبار القدرات", 24, "D96F63", bold=True, align="center", box_id=102)
# Subtitle: x=565.5, y=251.8, w=632.0, h=42.0
add_floating_textbox(doc, 565.5, 255.0, 632.0, 50.0, "(كمي ولفظي)", 18, "3BA4A6", bold=True, align="center", box_id=103)
# Badge: x=804.4, y=313.0, w=154.2, h=45.8
add_floating_textbox(doc, 804.4, 313.0, 154.2, 45.8, "دورة شاملة", 14, "FFFFFF", bold=True, align="center", box_id=104)

out_docx = os.path.abspath('test_tb_render.docx')
doc.save(out_docx)
print('Saved test_tb_render.docx')

# Render to PDF and image
word = win32.Dispatch('Word.Application')
word.Visible = False
try:
    wdoc = word.Documents.Open(out_docx)
    pdf_out = os.path.abspath('test_tb_render.pdf')
    wdoc.SaveAs2(pdf_out, FileFormat=17)
    wdoc.Close(False)
finally:
    try:
        word.Quit()
    except:
        pass

doc_pdf = pymupdf.open('test_tb_render.pdf')
for i, page in enumerate(doc_pdf):
    pix = page.get_pixmap(dpi=150)
    pix.save(f'test_tb_render_p{i+1}.png')
print('Rendered test_tb_render_p1.png')
