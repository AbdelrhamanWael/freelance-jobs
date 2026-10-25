import docx
from docx.shared import Inches, Pt
from docx.oxml import parse_xml
import json
import os
import win32com.client as win32

# Load JSON
with open('poster_texts.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

poster_box = data['poster_size']
texts = data['texts']

doc = docx.Document()
section = doc.sections[0]

# Set page size (convert px to inches. 96 dpi is standard web)
page_w_in = poster_box['width'] / 96.0
page_h_in = poster_box['height'] / 96.0
section.page_width = Inches(page_w_in)
section.page_height = Inches(page_h_in)
section.top_margin = Inches(0)
section.bottom_margin = Inches(0)
section.left_margin = Inches(0)
section.right_margin = Inches(0)

# Add background image
header = section.header
p_hdr = header.paragraphs[0]
run_hdr = p_hdr.add_run()
pic = run_hdr.add_picture('poster_bg_clean.png', width=Inches(page_w_in), height=Inches(page_h_in))

inline = run_hdr._r.xpath('.//wp:inline')[0]
extent = inline.xpath('.//wp:extent')[0]
cx = extent.get('cx')
cy = extent.get('cy')
graphic = inline.xpath('.//a:graphic')[0]

anchor_xml = f'''
<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
           xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
           distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
    <wp:simplePos x="0" y="0"/>
    <wp:positionH relativeFrom="page">
        <wp:posOffset>0</wp:posOffset>
    </wp:positionH>
    <wp:positionV relativeFrom="page">
        <wp:posOffset>0</wp:posOffset>
    </wp:positionV>
    <wp:extent cx="{cx}" cy="{cy}"/>
    <wp:effectExtent l="0" t="0" r="0" b="0"/>
    <wp:wrapNone/>
    <wp:docPr id="1" name="Background"/>
    <wp:cNvGraphicFramePr/>
</wp:anchor>
'''
anchor_el = parse_xml(anchor_xml)
anchor_el.append(graphic)
inline.getparent().replace(inline, anchor_el)

# PX to EMU for 96 DPI: 1 inch = 914400 EMUs, 1 px = 914400 / 96 = 9525 EMUs
PX_TO_EMU = 9525

def hex_to_w_val(rgb_str):
    if rgb_str.startswith('rgb'):
        vals = [int(v.strip()) for v in rgb_str.replace('rgb(', '').replace('rgba(', '').split(')')[0].split(',')]
        if len(vals) >= 3:
            return f"{vals[0]:02X}{vals[1]:02X}{vals[2]:02X}"
    return "000000"

def add_floating_textbox(doc, x_px, y_px, w_px, h_px, text, font_size_px, color, align="center", box_id=10):
    x_emu = int(x_px * PX_TO_EMU)
    y_emu = int(y_px * PX_TO_EMU)
    w_emu = int(w_px * PX_TO_EMU)
    h_emu = int(h_px * PX_TO_EMU)
    
    # font size in half points
    fs_px = float(font_size_px.replace('px', ''))
    # convert px to pt (px * 72 / 96)
    fs_pt = fs_px * 0.75
    sz_half = int(fs_pt * 2)
    
    color_hex = hex_to_w_val(color)
    
    if align == 'start' or align == 'right':
        align_w = 'right'
    elif align == 'left':
        align_w = 'left'
    else:
        align_w = 'center'
        
    tb_xml = f'''
    <w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
         xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
         xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
         xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
      <w:pPr>
        <w:spacing w:before="0" w:after="0" w:line="0" w:lineRule="auto"/>
      </w:pPr>
      <w:r>
        <w:drawing>
          <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251658240" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
            <wp:simplePos x="0" y="0"/>
            <wp:positionH relativeFrom="page">
              <wp:posOffset>{x_emu}</wp:posOffset>
            </wp:positionH>
            <wp:positionV relativeFrom="page">
              <wp:posOffset>{y_emu}</wp:posOffset>
            </wp:positionV>
            <wp:extent cx="{w_emu}" cy="{h_emu}"/>
            <wp:effectExtent l="0" t="0" r="0" b="0"/>
            <wp:wrapNone/>
            <wp:docPr id="{box_id}" name="TextBox_{box_id}"/>
            <wp:cNvGraphicFramePr/>
            <a:graphic>
              <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
                <wps:wsp>
                  <wps:spPr>
                    <a:xfrm>
                      <a:off x="0" y="0"/>
                      <a:ext cx="{w_emu}" cy="{h_emu}"/>
                    </a:xfrm>
                    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
                    <a:noFill/>
                    <a:ln><a:noFill/></a:ln>
                  </wps:spPr>
                  <wps:txbx>
                    <w:txbxContent>
                      <w:p>
                        <w:pPr>
                          <w:bidi/>
                          <w:jc w:val="{align_w}"/>
                          <w:spacing w:before="0" w:after="0"/>
                        </w:pPr>
                        <w:r>
                          <w:rPr>
                            <w:rtl/>
                            <w:rFonts w:ascii="Cairo" w:hAnsi="Cairo" w:cs="Cairo"/>
                            <w:b/>
                            <w:bCs/>
                            <w:color w:val="{color_hex}"/>
                            <w:sz w:val="{sz_half}"/>
                            <w:szCs w:val="{sz_half}"/>
                          </w:rPr>
                          <w:t>{text}</w:t>
                        </w:r>
                      </w:p>
                    </w:txbxContent>
                  </wps:txbx>
                  <wps:bodyPr lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr" wrap="none"/>
                </wps:wsp>
              </a:graphicData>
            </a:graphic>
          </wp:anchor>
        </w:drawing>
      </w:r>
    </w:p>
    '''
    try:
        p_el = parse_xml(tb_xml)
        doc._body._body.append(p_el)
    except Exception as e:
        print("XML Parsing error:", e)

box_id = 100
for item in texts:
    add_floating_textbox(
        doc=doc,
        x_px=item['x'],
        y_px=item['y'],
        w_px=item['width'] + 20, # Add a small buffer to avoid Word text wrapping
        h_px=item['height'] + 10,
        text=item['text'],
        font_size_px=item['fontSize'],
        color=item['color'],
        align=item['textAlign'],
        box_id=box_id
    )
    box_id += 1

out_docx = os.path.abspath('الإعلان_للتعديل_المريح.docx')
doc.save(out_docx)
print("Saved python-docx file.")

# Fix via COM
word = win32.Dispatch('Word.Application')
word.Visible = False
try:
    wdoc = word.Documents.Open(out_docx)
    final_docx = os.path.abspath('نسخة_العميل_قابلة_للتعديل_تماما.docx')
    wdoc.SaveAs2(final_docx, FileFormat=16)
    wdoc.Close(False)
    print("Fixed and saved final clean DOCX via Word COM.")
except Exception as e:
    print('COM Error:', e)
finally:
    try:
        word.Quit()
    except:
        pass
