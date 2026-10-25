import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client as win32
import os
import pymupdf

doc = docx.Document()
section = doc.sections[0]

# Page dimensions matching 1763 x 1377 aspect ratio (approx 11.5 x 8.98 inches)
page_w = Inches(11.5)
page_h = Inches(8.98)
section.page_width = page_w
section.page_height = page_h
section.top_margin = Inches(0.5)
section.bottom_margin = Inches(0.5)
section.left_margin = Inches(0.5)
section.right_margin = Inches(0.5)

header = section.header
p_hdr = header.paragraphs[0]
run_hdr = p_hdr.add_run()
pic = run_hdr.add_picture('test_bg_template_rendered.png', width=page_w, height=page_h)

# Convert inline picture to behindDoc anchor
inline = run_hdr._r.xpath('.//wp:inline')[0]
extent = inline.xpath('.//wp:extent')[0]
cx = extent.get('cx')
cy = extent.get('cy')
graphic = inline.xpath('.//a:graphic')[0]

anchor_xml = f'''
<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
           xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
           distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
    <wp:simplePos x="0" y="0"/>
    <wp:positionH relativeFrom="page">
        <wp:posOffset>0</wp:posOffset>
    </wp:positionH>
    <wp:positionV relativeFrom="page">
        <wp:posOffset>0</wp:posOffset>
    </wp:positionV>
    <wp:extent cx="{cx}" cy="{cy}"/>
    <wp:effectExtent l="0" t="0" r="0" b="0"/>
    <wp:wrapNone/>
    <wp:docPr id="1" name="Background"/>
    <wp:cNvGraphicFramePr/>
</wp:anchor>
'''
anchor_el = parse_xml(anchor_xml)
anchor_el.append(graphic)
inline.getparent().replace(inline, anchor_el)

# Add sample body text
p = doc.add_paragraph()
p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('معهد مئة وثمانون درجة العالي للتدريب')
run.font.size = Pt(24)
run.font.bold = True
run.font.color.rgb = RGBColor(217, 111, 99)

test_docx = os.path.abspath('test_anchor.docx')
doc.save(test_docx)
print('Saved test_anchor.docx')

word = win32.Dispatch('Word.Application')
word.Visible = False
try:
    wdoc = word.Documents.Open(test_docx)
    pdf_out = os.path.abspath('test_anchor.pdf')
    wdoc.SaveAs2(pdf_out, FileFormat=17)
    wdoc.Close(False)
finally:
    try:
        word.Quit()
    except:
        pass

doc_pdf = pymupdf.open('test_anchor.pdf')
print('PDF pages:', len(doc_pdf))
for i, page in enumerate(doc_pdf):
    pix = page.get_pixmap(dpi=150)
    pix.save(f'test_anchor_p{i+1}.png')
print('Rendered test_anchor preview!')
