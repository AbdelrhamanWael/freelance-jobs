import openpyxl

# Load the Excel file
wb = openpyxl.load_workbook('FLEET_JOBS_WITH_EASY_APPLY_COLUMN.xlsx')
ws = wb.active

# Get all jobs (skip header rows)
jobs = []
for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
    title = row[1].value
    company = row[2].value
    if title and company:
        jobs.append({'title': title, 'company': company, 'row': row[0].row})

# Check for duplicates
seen = {}
duplicates = []
for job in jobs:
    key = f"{job['title']}|{job['company']}"
    if key in seen:
        duplicates.append(job)
        print(f"DUPLICATE FOUND: {job['title']} @ {job['company']} (Row {job['row']})")
        print(f"  Original at: Row {seen[key]}")
    else:
        seen[key] = job['row']

if not duplicates:
    print("No duplicates found! File is clean.")
else:
    print(f"\nTotal duplicates found: {len(duplicates)}")
