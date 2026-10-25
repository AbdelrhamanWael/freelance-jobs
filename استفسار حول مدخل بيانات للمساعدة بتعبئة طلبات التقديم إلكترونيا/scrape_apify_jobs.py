import os
import sys
import time
import requests
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Ensure UTF-8 output on Windows console
sys.stdout.reconfigure(encoding='utf-8')

TOKEN = os.environ.get("APIFY_TOKEN", "")
ACTOR = "agentx~all-jobs-scraper"  # Apify actor for multi-platform job scraping

# 2. Target Keywords (Curated for Mechanical / Fleet / Heavy Equipment Engineer - Ahmed Aboud)
SEARCH_QUERIES = [
    # Top Priority Keywords from Strategy Document
    "Fleet Maintenance Engineer",
    "Fleet Maintenance Supervisor",
    "Maintenance Workshop Supervisor",
    "Maintenance & Reliability Engineer",
    "Heavy Equipment Maintenance Engineer",
    "Fleet Operations & Maintenance Engineer",
    "Bus Maintenance Engineer",
    
    # Secondary & Supporting Keywords
    "Heavy Equipment Engineer",
    "Transport Coordinator",
    "Fleet Maintenance Manager",
    "Workshop Supervisor",
    "Vehicle Maintenance Engineer"
]

# Target Country & Locations
COUNTRY = "Saudi Arabia"
MAX_RESULTS_PER_KEYWORD = 30  # Adjust as needed (e.g. 20 - 50)


def scrape_all_jobs(output_dir="."):
    print("=" * 70)
    print("🚀 Starting Job Scraping via Apify (agentx~all-jobs-scraper)")
    print(f"📍 Country: {COUNTRY}")
    print(f"🔑 Keywords count: {len(SEARCH_QUERIES)}")
    print("=" * 70)

    all_jobs = []
    seen_urls = set()
    seen_signatures = set()

    for idx, kw in enumerate(SEARCH_QUERIES, 1):
        print(f"\n[{idx}/{len(SEARCH_QUERIES)}] 🔍 Searching keyword: '{kw}'...")
        payload = {
            "keyword": kw,
            "country": COUNTRY,
            "max_results": MAX_RESULTS_PER_KEYWORD
        }

        try:
            r = requests.post(
                f"https://api.apify.com/v2/acts/{ACTOR}/run-sync-get-dataset-items",
                params={"token": TOKEN},
                json=payload,
                timeout=240
            )

            if r.status_code not in [200, 201]:
                print(f"   ⚠️ API returned status {r.status_code}: {r.text[:200]}")
                continue

            raw_items = r.json()
            if not isinstance(raw_items, list):
                print(f"   ⚠️ Unexpected response format: {str(raw_items)[:200]}")
                continue

            print(f"   ✅ Raw items retrieved: {len(raw_items)}")
            kw_valid_count = 0

            for item in raw_items:
                platform = item.get("platform") or "Web"
                # Exclude freelance / gig tasks
                if platform.lower() in ["freelancer.com", "upwork", "fiverr"]:
                    continue

                title = (item.get("title") or item.get("jobTitle") or "").strip()
                company = (item.get("company_name") or item.get("company") or "").strip()
                url = (item.get("official_url") or item.get("platform_url") or item.get("url") or "").strip()

                if not title or not url:
                    continue

                # Format location
                loc_data = item.get("location")
                if isinstance(loc_data, dict):
                    locality = loc_data.get("locality") or ""
                    region = loc_data.get("region") or ""
                    raw_loc = loc_data.get("raw") or ""
                    country_name = loc_data.get("country") or "Saudi Arabia"
                    
                    parts = [p for p in [locality, region, country_name] if p]
                    loc_str = ", ".join(parts) if parts else raw_loc or "Saudi Arabia"
                elif isinstance(loc_data, str) and loc_data:
                    loc_str = loc_data
                else:
                    loc_str = "Saudi Arabia"

                # Deduplication check
                signature = f"{title.lower()}::{company.lower()}"
                if url in seen_urls or signature in seen_signatures:
                    continue

                seen_urls.add(url)
                seen_signatures.add(signature)

                job_record = {
                    "title": title,
                    "company": company or "Confidential",
                    "location": loc_str,
                    "platform": platform,
                    "url": url,
                    "posted_date": item.get("posted_date") or item.get("postingDate") or "",
                    "work_mode": item.get("work_mode") or "On-site",
                    "matched_keyword": kw
                }

                all_jobs.append(job_record)
                kw_valid_count += 1

            print(f"   🎯 Added {kw_valid_count} new unique jobs from this keyword.")

        except requests.exceptions.Timeout:
            print(f"   ⏱️ Request timed out for keyword '{kw}'. Skipping to next.")
        except Exception as e:
            print(f"   ❌ Error searching '{kw}': {e}")

        # Rate-limiting pause
        time.sleep(1)

    print("\n" + "=" * 70)
    print(f"🎉 Scraping Complete! Total unique jobs gathered: {len(all_jobs)}")
    print("=" * 70)

    # Save to JSON
    json_path = os.path.join(output_dir, "apify_scraped_jobs.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(all_jobs, f, ensure_ascii=False, indent=2)
    print(f"💾 Saved JSON results to: {json_path}")

    # Save to Styled Excel
    excel_path = os.path.join(output_dir, "Apify_Scraped_Jobs_Saudi.xlsx")
    export_to_excel(all_jobs, excel_path)
    print(f"📊 Saved Styled Excel results to: {excel_path}")


def export_to_excel(jobs, file_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scraped Jobs"

    # Header Row
    headers = [
        "#", "Job Title", "Company", "Location", "Platform",
        "Work Mode", "Posted Date", "Search Keyword", "Job Link"
    ]
    ws.row_dimensions[1].height = 28

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col)
        cell.value = h
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Styling for data rows
    regular_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")
    
    fill_even = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_odd = PatternFill(fill_type=None)
    
    thin_side = Side(border_style="thin", color="D9D9D9")
    data_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for i, j in enumerate(jobs, 1):
        row_num = i + 1
        ws.row_dimensions[row_num].height = 24
        row_fill = fill_even if i % 2 == 0 else fill_odd

        row_data = [
            i,
            j["title"],
            j["company"],
            j["location"],
            j["platform"],
            j["work_mode"],
            j["posted_date"][:10] if j["posted_date"] else "Recent",
            j["matched_keyword"],
            "View Job"
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row_num, col_idx)
            cell.value = val
            cell.fill = row_fill
            cell.border = data_border
            
            if col_idx in [1, 5, 6, 7, 8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            if col_idx == 2:
                cell.font = bold_font
            elif col_idx == 9: # Link column
                cell.font = link_font
                cell.hyperlink = j["url"]
            else:
                cell.font = regular_font

    # Column Widths
    col_widths = {
        "A": 6.0,   # #
        "B": 42.0,  # Job Title
        "C": 30.0,  # Company
        "D": 25.0,  # Location
        "E": 16.0,  # Platform
        "F": 14.0,  # Work Mode
        "G": 14.0,  # Posted Date
        "H": 32.0,  # Search Keyword
        "I": 16.0   # Job Link
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:I{len(jobs) + 1}"
    ws.views.sheetView[0].showGridLines = True
    wb.save(file_path)


if __name__ == "__main__":
    scrape_all_jobs()
