import os
import sys
if sys.stdout.encoding.lower() != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')
import json
import time
import re
import pandas as pd
import requests
import docx
from pydantic import BaseModel, Field, ValidationError

# ==========================================
# 1. CONFIGURATION
# ==========================================
MISTRAL_API_KEY = "Ab7xuVjRVDTILXKVbHDvJHwpC0bO4aYr"
EXCEL_FILE = "Combined_FB_Data.xlsx"
TEMPLATE_FILE = "استمارة-معدلة.docx"

MODEL_NAME = "mistral-large-latest"
TEST_MODE = False

MAX_RETRIES_PER_POST = 3
MAX_RATE_LIMIT_RETRIES = 5
BASE_SLEEP_BETWEEN_POSTS = 2
FAILED_RAW_DIR = "Failed_Raw"
# ==========================================

# ==========================================
# 2. DEFINE THE STRUCTURED OUTPUT SCHEMA
# ==========================================
class ContextualAnalysis(BaseModel):
    occasion: str = Field(description="المناسبة أو الحدث (استنتجها من النص)")
    communication_goal: str = Field(description="الهدف الاتصالي الظاهر (اعلام/ اقناع/ حوار/ بناء اجماع)")
    target_audience: str = Field(description="الجمهور المستهدف (الجاليات/ جمهور دولي/ جمهور إحدى الدول/ صناع قرار/ وسائل اعلام/ اخرى)")
    actors: str = Field(description="الأطراف الفاعلة (رئيس الدولة/ وزير/ دبلوماسي/ منظمة دولية/ سفارة/ اخرى)")
    regional_context: str = Field(description="السياق الإقليمي والدولي (توضيح السياق في سطرين)")
    geographic_scope: str = Field(description="النطاق الجغرافي (إقليمي / دولي)")
    qualitative_analysis_space: str = Field(description="مساحة التحليل الكيفي (تفسير وتحليل عام للمنشور)")

class VanDijkModel(BaseModel):
    highlight_self_positives: Tactic = Field(description="إبراز إيجابيات الذات")
    hide_self_negatives: Tactic = Field(description="إخفاء سلبيات الذات")
    highlight_other_negatives: Tactic = Field(description="إبراز سلبيات الآخر")
    hide_other_positives: Tactic = Field(description="إخفاء إيجابيات الآخر")

class InfoStrategy(BaseModel):
    official_info: Tactic = Field(description="نشر معلومات أو بيانات رسمية")
    officials_statements: Tactic = Field(description="نشر تصريحات المسؤولين")
    intl_coop: Tactic = Field(description="عرض مبادرات أو علاقات دولية (تعاون دولي)")
    diplomatic_visits: Tactic = Field(description="الإعلان عن اللقاءات والزيارات الدبلوماسية")
    stats_reports: Tactic = Field(description="نشر التقارير والاحصاءات")

class PersuasionStrategy(BaseModel):
    intl_sources: Tactic = Field(description="الاستشهاد بمصادر دولية")
    human_stories: Tactic = Field(description="ابراز قصص انسانية")
    national_success: Tactic = Field(description="إبراز النجاحات الوطنية")
    state_superiority: Tactic = Field(description="ابراز تفوق او مكانة الدولة")
    solidarity: Tactic = Field(description="التهنئة بالمناسبات او التضامن")

class ReputationStrategy(BaseModel):
    political_values: Tactic = Field(description="إبراز القيم السياسية أو الإنسانية")
    tech_progress: Tactic = Field(description="إبراز التقدم العلمي والتكنولوجي")
    intl_leadership: Tactic = Field(description="إبراز الدور القيادي الدولي")
    reliable_partner: Tactic = Field(description="تقديم الدولة كشريك موثوق")
    correct_stereotypes: Tactic = Field(description="تصحيح الصور النمطية السلبية")
    intl_trust: Tactic = Field(description="بناء الثقة الدولية")
    cultural_appeal: Tactic = Field(description="إبراز الجاذبية الحضارية والثقافية")

class InteractionStrategy(BaseModel):
    hashtags: Tactic = Field(description="استخدام الوسوم (Hashtags)")
    mentions: Tactic = Field(description="استخدام الإشارات (@Mention)")
    retweets: Tactic = Field(description="إعادة نشر أو اقتباس التغريدات")
    call_to_action: Tactic = Field(description="دعوة الجمهور للمشاركة")
    official_links: Tactic = Field(description="توظيف روابط لمصادر رسمية")

class DialogueStrategy(BaseModel):
    comms_channels: Tactic = Field(description="توفير قنوات للتواصل وطلب المعلومات")
    reply_audience: Tactic = Field(description="الرد على تعليقات واستفسارات الجمهور")
    live_sessions: Tactic = Field(description="تنظيم جلسات أو لقاءات للحوار المباشر")
    encourage_proposals: Tactic = Field(description="تشجيع الجمهور على تقديم المقترحات")
    polls_questions: Tactic = Field(description="طرح الأسئلة واستطلاعات الرأي العام")

class DigitalDiplomacyStrategies(BaseModel):
    information_strategy: InfoStrategy
    persuasion_strategy: PersuasionStrategy
    reputation_strategy: ReputationStrategy
    interaction_strategy: InteractionStrategy
    dialogue_strategy: DialogueStrategy

class Tactic(BaseModel):
    present: bool
    notes: str = Field(description="تحليل مفصل وعميق إجباري (يمنع تركه فارغاً أو كتابة نقاط. اشرح السبب سواء كانت موجودة أم غائبة)")

class EmotionalTactics(BaseModel):
    irony: Tactic = Field(description="السخرية")
    celeb_quotes: Tactic = Field(description="الاستشهاد بالشخصيات المشهورة")
    disapproval: Tactic = Field(description="الالاستهجان")
    politeness: Tactic = Field(description="التلطف")
    show_power: Tactic = Field(description="استعراض القوة")
    cultural_symbols: Tactic = Field(description="استخدام الرموز الثقافية المألوفة")
    national_praise: Tactic = Field(description="تمجيد الذات القومية")
    endearment: Tactic = Field(description="التحبيب والاستحسان")
    doubt: Tactic = Field(description="التشكيك والتكذيب")
    linguistic_styles: Tactic = Field(description="الأساليب اللغوية")
    threats: Tactic = Field(description="التخويف والتهديد")
    victimhood: Tactic = Field(description="لعب دور الضحية")
    similarities: Tactic = Field(description="التأكيد على أوجه الشبه والتقارب")
    challenge: Tactic = Field(description="التحدي")
    slogans: Tactic = Field(description="الشعارات")
    citizen_opinions: Tactic = Field(description="الاستشهاد بآراء المواطنين")
    categorization: Tactic = Field(description="التصنيف")
    positive_traits: Tactic = Field(description="الصفات الايجابية")
    other_negatives_show: Tactic = Field(description="عرض سلبيات الآخر")
    superlatives: Tactic = Field(description="استخدام صيغة التفضيل")
    supportive_power: Tactic = Field(description="القوة الداعمة")

class LogicalTactics(BaseModel):
    examples: Tactic = Field(description="التوضيح بالامثلة")
    comparison: Tactic = Field(description="المقارنة")
    proof: Tactic = Field(description="البرهنة")
    justification: Tactic = Field(description="التبرير")
    stats_numbers: Tactic = Field(description="الأرقام والاحصاءات")
    number_game: Tactic = Field(description="لعبة الارقام")
    religious_frames: Tactic = Field(description="توظيف الأطر المرجعية الدينية")

class SoftPowerItem(BaseModel):
    presence: str = Field(description="اختر واحدة فقط: 'رئيسي' أو 'ثانوي' أو 'غائب'")
    notes: str

class SoftPowerDynamics(BaseModel):
    political: SoftPowerItem
    cultural: SoftPowerItem
    educational: SoftPowerItem
    economic: SoftPowerItem
    scientific: SoftPowerItem
    humanitarian: SoftPowerItem
    media: SoftPowerItem

class PostAnalysis(BaseModel):
    contextual_analysis: ContextualAnalysis
    van_dijk_model: VanDijkModel
    digital_diplomacy_strategies: DigitalDiplomacyStrategies
    emotional_tactics: EmotionalTactics
    logical_tactics: LogicalTactics
    soft_power_dynamics: SoftPowerDynamics

# ==========================================
# 3. HELPER FUNCTIONS
# ==========================================

def _strip_code_fences(text: str) -> str:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```[a-zA-Z]*\n?", "", text)
    if text.endswith("```"):
        text = text[: -3]
    return text.strip()

def _extract_json_block(text: str) -> str:
    text = _strip_code_fences(text)
    start = text.find("{")
    end = text.rfind("}")
    if start == -1 or end == -1 or end <= start:
        raise ValueError("No JSON object found in model response")
    return text[start : end + 1]

def _remove_control_characters(text: str) -> str:
    return re.sub(r"[\x00-\x1f\x7f]", " ", text)

def _unwrap_known_wrapper(data: dict) -> dict:
    expected_keys = set(PostAnalysis.model_fields.keys())
    if expected_keys.issubset(data.keys()):
        return data
    if len(data) == 1:
        inner = next(iter(data.values()))
        if isinstance(inner, dict) and expected_keys.issubset(inner.keys()):
            return inner
    return data

def parse_ai_response(raw_text: str) -> PostAnalysis:
    json_str = _extract_json_block(raw_text)
    json_str = _remove_control_characters(json_str)
    data = json.loads(json_str)
    data = _unwrap_known_wrapper(data)
    return PostAnalysis.model_validate(data)

def call_mistral_api(post_text: str, metrics: str) -> str:
    system_prompt = (
        "أنت باحث أكاديمي خبير في 'تحليل الخطاب الكيفي' والدبلوماسية الرقمية.\n"
        "مهمتك هي تحليل المنشور التالي بدقة وبشكل شامل وتفصيلي وتعبئة الاستمارة المطلوبة بالكامل.\n"
        "تعليمات هامة جداً (إجبارية):\n"
        "1. يمنع منعاً باتاً ترك أي حقل 'ملاحظات' (notes) فارغاً أو كتابة (...) أو (لا يوجد). يجب تعبئة كل الحقول بدون استثناء.\n"
        "2. لكل استراتيجية أو تكتيك، إذا اخترت True (نعم)، اشرح كيف تم استخدامها مع ذكر أمثلة من النص. وإذا اخترت False (لا)، يجب أن تكتب تحليلاً يشرح لماذا غابت هذه الاستراتيجية وتأثير غيابها على المنشور.\n"
        "3. التحليل الكيفي مهم جداً: لا تكتفِ بالحديث عن 'نعم' أو 'لا' في حقول الملاحظات، بل قدم تحليلاً وافياً وكافياً وعميقاً.\n"
        "4. الاتساق المنطقي التام: يمنع منعاً باتاً التناقض بين الخيارات المنطقية والتحليل المكتوب.\n"
        "أعد الناتج بصيغة JSON صرفة فقط، مطابقة تماماً للمخطط المطلوب."
    )

    user_prompt = (
        f"قم بتحليل هذا المنشور:\n\nنص المنشور:\n{post_text}\n\nبيانات تفاعلية:\n{metrics}\n\n"
        f"أعد الناتج بصيغة JSON فقط متوافق مع المخطط التالي:\n{PostAnalysis.model_json_schema()}"
    )

    headers = {
        "Authorization": f"Bearer {MISTRAL_API_KEY}",
        "Content-Type": "application/json",
    }

    payload = {
        "model": MODEL_NAME,
        "temperature": 0.1,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "response_format": {"type": "json_object"},
    }

    for attempt in range(1, MAX_RATE_LIMIT_RETRIES + 1):
        response = requests.post("https://api.mistral.ai/v1/chat/completions", headers=headers, json=payload, timeout=300)
        if response.status_code == 429:
            time.sleep(10)
            continue
        response.raise_for_status()
        return response.json()["choices"][0]["message"]["content"]

    raise RuntimeError("Exceeded max retries for rate limiting (429)")

def analyze_post_with_ai(post_text: str, metrics: str, debug_tag: str) -> PostAnalysis:
    last_error = None
    last_raw = None
    for attempt in range(1, MAX_RETRIES_PER_POST + 1):
        raw_text = call_mistral_api(post_text, metrics)
        last_raw = raw_text
        try:
            return parse_ai_response(raw_text)
        except Exception as e:
            last_error = e
            print(f"    ⚠️ Parse attempt {attempt}/{MAX_RETRIES_PER_POST} failed: {e}")
            time.sleep(2)
            
    os.makedirs(FAILED_RAW_DIR, exist_ok=True)
    dump_path = os.path.join(FAILED_RAW_DIR, f"{debug_tag}_raw.txt")
    with open(dump_path, "w", encoding="utf-8") as f: f.write(last_raw or "")
    raise RuntimeError(f"Failed to parse JSON. Saved raw to {dump_path}: {last_error}")

def yn(val: bool) -> str:
    return "☑ نعم    ☐ لا" if val else "☐ نعم    ☑ لا"

def generate_word_report(analysis: PostAnalysis, post_data: dict, output_filename: str):
    """Fills the template DOCX directly to match exact original format."""
    doc = docx.Document(TEMPLATE_FILE)

    post_text = str(post_data.get('Post Text', ''))

    # Update Platform Name (Table 1)
    if len(doc.tables) > 1 and len(doc.tables[1].rows) > 0 and len(doc.tables[1].rows[0].cells) > 2:
        doc.tables[1].rows[0].cells[2].text = "منصة فيسبوك (Facebook)"
    
    # Clear old Sample Accounts (Table 2)
    post_url = post_data.get("Post URL", "")
    if len(doc.tables) > 2 and len(doc.tables[2].rows) > 1:
        if "USEmbassyCairo" in post_url:
            doc.tables[2].rows[1].cells[0].text = "-"
            doc.tables[2].rows[1].cells[1].text = "U.S. Embassy Cairo (سفارة أمريكا بالقاهرة)"
        elif "ChineseEmbinEgypt" in post_url:
            doc.tables[2].rows[1].cells[0].text = "Chinese Embassy in Egypt (سفارة الصين بالقاهرة)"
            doc.tables[2].rows[1].cells[1].text = "-"
        else:
            doc.tables[2].rows[1].cells[0].text = "حساب صيني"
            doc.tables[2].rows[1].cells[1].text = "حساب أمريكي"

    # Metadata (Table 4)
    tbl4 = doc.tables[4]
    likes = post_data.get('Likes/Reactions', '0')
    comments = post_data.get('Comments', '0')
    shares = post_data.get('Shares', '0')
    media_type = post_data.get('Media Type', 'N/A')
    
    tbl4.rows[1].cells[0].text = f"إعجابات: {likes}، تعليقات: {comments}، مشاركات: {shares}"
    tbl4.rows[1].cells[1].text = str(media_type)
    tbl4.rows[1].cells[2].text = analysis.contextual_analysis.occasion
    tbl4.rows[1].cells[3].text = "العربية"
    tbl4.rows[1].cells[4].text = str(post_data.get('Post URL', ''))
    tbl4.rows[1].cells[6].text = str(post_data.get('Account', ''))
    tbl4.rows[1].cells[7].text = str(post_data.get('Publication Date & Time', ''))
    tbl4.rows[1].cells[8].text = "1"
    
    # Remove empty rows 2 to 16 in Table 4
    for _ in range(15):
        if len(tbl4.rows) > 2: tbl4._element.remove(tbl4.rows[2]._element)

    # Contextual (Table 5)
    tbl5 = doc.tables[5]
    c = analysis.contextual_analysis
    tbl5.rows[2].cells[0].text = c.qualitative_analysis_space
    tbl5.rows[2].cells[2].text = c.communication_goal
    tbl5.rows[2].cells[3].text = c.target_audience
    tbl5.rows[2].cells[4].text = c.actors
    tbl5.rows[2].cells[5].text = c.regional_context
    tbl5.rows[2].cells[6].text = c.geographic_scope
    tbl5.rows[2].cells[7].text = (post_text[:200] + "...") if len(post_text) > 200 else post_text
    tbl5.rows[2].cells[8].text = "1"
    
    # Remove empty rows 3 to 16 in Table 5
    for _ in range(14):
        if len(tbl5.rows) > 3: tbl5._element.remove(tbl5.rows[3]._element)

    # Van Dijk (Table 6)
    tbl6 = doc.tables[6]
    v = analysis.van_dijk_model
    tbl6.rows[3].cells[0].text = ("☑" if v.highlight_self_positives.present else "☐") + " - " + v.highlight_self_positives.notes
    tbl6.rows[4].cells[0].text = ("☑" if v.hide_self_negatives.present else "☐") + " - " + v.hide_self_negatives.notes
    tbl6.rows[5].cells[0].text = ("☑" if v.highlight_other_negatives.present else "☐") + " - " + v.highlight_other_negatives.notes
    tbl6.rows[6].cells[0].text = ("☑" if v.hide_other_positives.present else "☐") + " - " + v.hide_other_positives.notes

    # Digital Diplomacy (Table 7)
    tbl7 = doc.tables[7]
    
    def fill_strat(row_idx, tactic):
        tbl7.rows[row_idx].cells[1].text = yn(tactic.present)
        tbl7.rows[row_idx].cells[0].text = tactic.notes

    info = analysis.digital_diplomacy_strategies.information_strategy
    fill_strat(3, info.official_info)
    fill_strat(4, info.officials_statements)
    fill_strat(5, info.intl_coop)
    fill_strat(6, info.diplomatic_visits)
    fill_strat(7, info.stats_reports)

    pers = analysis.digital_diplomacy_strategies.persuasion_strategy
    fill_strat(10, pers.intl_sources)
    fill_strat(11, pers.human_stories)
    fill_strat(12, pers.national_success)
    fill_strat(13, pers.state_superiority)
    fill_strat(14, pers.solidarity)

    rep = analysis.digital_diplomacy_strategies.reputation_strategy
    fill_strat(17, rep.political_values)
    fill_strat(18, rep.tech_progress)
    fill_strat(19, rep.intl_leadership)
    fill_strat(20, rep.reliable_partner)
    fill_strat(21, rep.correct_stereotypes)
    fill_strat(22, rep.intl_trust)
    fill_strat(23, rep.cultural_appeal)

    inte = analysis.digital_diplomacy_strategies.interaction_strategy
    fill_strat(26, inte.hashtags)
    fill_strat(27, inte.mentions)
    fill_strat(28, inte.retweets)
    fill_strat(29, inte.call_to_action)
    fill_strat(30, inte.official_links)

    diag = analysis.digital_diplomacy_strategies.dialogue_strategy
    # Dialogue starts at row 33 in the template
    fill_strat(33, diag.comms_channels)
    fill_strat(34, diag.reply_audience)
    fill_strat(35, diag.live_sessions)
    fill_strat(36, diag.encourage_proposals)
    fill_strat(37, diag.polls_questions)

    # Tactics (Table 8)
    tbl8 = doc.tables[8]
    def fill_tactic(row_idx, tactic):
        tbl8.rows[row_idx].cells[1].text = "☑" if tactic.present else "☐"
        tbl8.rows[row_idx].cells[2].text = "☐" if tactic.present else "☑"
        tbl8.rows[row_idx].cells[5].text = tactic.notes

    emot = analysis.emotional_tactics
    fill_tactic(3, emot.irony)
    fill_tactic(4, emot.celeb_quotes)
    fill_tactic(5, emot.disapproval)
    fill_tactic(6, emot.politeness)
    fill_tactic(7, emot.show_power)
    fill_tactic(8, emot.cultural_symbols)
    fill_tactic(9, emot.national_praise)
    fill_tactic(10, emot.endearment)
    fill_tactic(11, emot.doubt)
    fill_tactic(12, emot.linguistic_styles)
    fill_tactic(13, emot.threats)
    fill_tactic(14, emot.victimhood)
    fill_tactic(15, emot.similarities)
    fill_tactic(16, emot.challenge)
    fill_tactic(17, emot.slogans)
    fill_tactic(18, emot.citizen_opinions)
    fill_tactic(19, emot.categorization)
    fill_tactic(20, emot.positive_traits)
    fill_tactic(21, emot.other_negatives_show)
    fill_tactic(22, emot.superlatives)
    fill_tactic(23, emot.supportive_power)

    log = analysis.logical_tactics
    fill_tactic(26, log.examples)
    fill_tactic(27, log.comparison)
    fill_tactic(28, log.proof)
    fill_tactic(29, log.justification)
    fill_tactic(30, log.stats_numbers)
    fill_tactic(31, log.number_game)
    fill_tactic(32, log.religious_frames)

    # Soft Power (Table 9)
    tbl9 = doc.tables[9]
    def sp_text(val):
        return f"{'☑' if 'رئيسي' in val else '☐'} رئيسي    {'☑' if 'ثانوي' in val else '☐'} ثانوي    {'☑' if 'غائب' in val else '☐'} غائب"
    
    sp = analysis.soft_power_dynamics
    tbl9.rows[2].cells[1].text = sp_text(sp.political.presence)
    tbl9.rows[2].cells[0].text = sp.political.notes
    tbl9.rows[3].cells[1].text = sp_text(sp.cultural.presence)
    tbl9.rows[3].cells[0].text = sp.cultural.notes
    tbl9.rows[4].cells[1].text = sp_text(sp.educational.presence)
    tbl9.rows[4].cells[0].text = sp.educational.notes
    tbl9.rows[5].cells[1].text = sp_text(sp.economic.presence)
    tbl9.rows[5].cells[0].text = sp.economic.notes
    tbl9.rows[6].cells[1].text = sp_text(sp.scientific.presence)
    tbl9.rows[6].cells[0].text = sp.scientific.notes
    tbl9.rows[7].cells[1].text = sp_text(sp.humanitarian.presence)
    tbl9.rows[7].cells[0].text = sp.humanitarian.notes
    tbl9.rows[8].cells[1].text = sp_text(sp.media.presence)
    tbl9.rows[8].cells[0].text = sp.media.notes

    doc.save(output_filename)

# ==========================================
# 4. MAIN EXECUTION LOOP
# ==========================================
def main():
    if not os.path.exists(EXCEL_FILE):
        print(f"❌ ERROR: Cannot find {EXCEL_FILE}.")
        return
    if not os.path.exists(TEMPLATE_FILE):
        print(f"❌ ERROR: Cannot find {TEMPLATE_FILE}.")
        return

    print("📊 Loading Excel data...")
    df = pd.read_excel(EXCEL_FILE)
    df = df.dropna(subset=["Post Text"])

    output_dir = "Analysis_Reports"
    os.makedirs(output_dir, exist_ok=True)

    limit = 3 if TEST_MODE else len(df)
    print(f"🚀 Starting analysis for {limit} posts (TEST_MODE={TEST_MODE}, model={MODEL_NAME})...")

    success_count = 0
    failed_indices = []

    for index, row in df.head(limit).iterrows():
        post_text = str(row["Post Text"])
        metrics = f"الإعجابات: {row.get('Likes/Reactions', '0')} | التعليقات: {row.get('Comments', '0')} | المشاركات: {row.get('Shares', '0')}"

        safe_index = str(index + 1).zfill(3)
        post_url = str(row.get("Post URL", ""))
        
        if "USEmbassyCairo" in post_url:
            acc_folder = "US_Embassy_Cairo"
        elif "ChineseEmbinEgypt" in post_url:
            acc_folder = "Chinese_Embassy_Egypt"
        else:
            acc_folder = "Other_Accounts"
            
        account_path = os.path.join(output_dir, acc_folder)
        os.makedirs(account_path, exist_ok=True)
        
        filename = os.path.join(account_path, f"Analysis_Post_{safe_index}.docx")
        
        if os.path.exists(filename):
            print(f"  ⏭️ Skipping Post #{index + 1}: Already analyzed.")
            success_count += 1
            continue
            
        print(f"⏳ Analyzing Post #{index + 1}...")
        try:
            analysis = analyze_post_with_ai(post_text, metrics, debug_tag=f"Post_{safe_index}")
            generate_word_report(analysis, row.to_dict(), filename)
            print(f"  ✅ Saved: {filename}")
            success_count += 1
        except Exception as e:
            print(f"  ❌ Failed Post #{index + 1}: {e}")
            failed_indices.append(index + 1)

        if index + 1 < limit:
            print(f"  💤 Waiting {BASE_SLEEP_BETWEEN_POSTS} seconds...")
            time.sleep(BASE_SLEEP_BETWEEN_POSTS)

    print(f"\n🎉 Finished! Successfully analyzed {success_count}/{limit} posts.")
    if TEST_MODE: print("ℹ️ Note: TEST_MODE is ON.")

if __name__ == "__main__":
    main()