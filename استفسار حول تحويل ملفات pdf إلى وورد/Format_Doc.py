import os
import sys
import re
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

# Fix for printing Arabic text in Windows console
if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')

def set_rtl(paragraph):
    pPr = paragraph._element.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)
    # Don't override alignment if it's already centered
    if paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def apply_font(run, bold=False, size=12, color="000000"):
    run.font.name = 'Arial'
    run._element.rPr.rFonts.set(qn('w:cs'), 'Arial')
    run.font.size = Pt(size)
    run.bold = bold
    if color != "000000" and color != "black":
        try:
            hex_color = color.lstrip('#')
            run.font.color.rgb = RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))
        except:
            pass

def set_cell_background(cell, fill="DDEBF7"):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def is_heading(text):
    # المادة الأولى، المادة الثانية، الوثائق المالية، الخ
    return text.startswith("المادة") or re.match(r'^\d+\.\s+.*', text) or text.startswith("نموذج") or "سياسة حفظ الوثائق" in text

def is_bullet(text):
    # السجلات المالية، اللائحة الأساسية، سجل، ملفات، إلخ (عادة تبدأ بنقاط في الملف الأصلي)
    return text.startswith("سجل") or text.startswith("ملف") or text.startswith("اللائحة") or text.startswith("دليل") or text.startswith("شؤون") or text.startswith("جميع")

def format_document(input_path, output_path):
    print("بدء تنسيق المستند وتنظيفه...")
    doc = Document(input_path)

    # 1. تنسيق الفقرات
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue

        # تنظيف علامات الـ Markdown (**)
        if "**" in text:
            text = text.replace("**", "")
            
            # إعادة كتابة النص لضمان إزالة النجوم من الـ runs
            if len(p.runs) > 0:
                p.runs[0].text = text
                for i in range(1, len(p.runs)):
                    p.runs[i].text = ""

        set_rtl(p)
        
        # التأكد من التباعد بين السطور
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.2

        if len(p.runs) > 0:
            run = p.runs[0]
            
            if is_heading(text):
                # تنسيق العناوين (أزرق، عريض، كبير)
                apply_font(run, bold=True, size=14, color="2E74B5")
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(6)
                if "سياسة حفظ الوثائق" in text or "نموذج" in text:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif is_bullet(text):
                # تنسيق النقاط
                apply_font(run, bold=False, size=12, color="000000")
                if not text.startswith("-") and not text.startswith("•"):
                    p.runs[0].text = "• " + p.runs[0].text
            else:
                # النص العادي
                apply_font(run, bold=False, size=12, color="000000")

    # 2. تنسيق الجداول
    for table in doc.tables:
        # محاذاة الجدول RTL
        tblPr = table._element.tblPr
        bidiVisual = OxmlElement('w:bidiVisual')
        tblPr.append(bidiVisual)

        for i, row in enumerate(table.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    text = p.text.replace("**", "")
                    if len(p.runs) > 0:
                        p.runs[0].text = text
                        for idx in range(1, len(p.runs)):
                            p.runs[idx].text = ""
                            
                    set_rtl(p)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    
                    if len(p.runs) > 0:
                        run = p.runs[0]
                        if i == 0:
                            # الصف الأول (الترويسة)
                            set_cell_background(cell, "DDEBF7") # أزرق فاتح جداً
                            apply_font(run, bold=True, size=12, color="000000")
                        else:
                            # باقي الصفوف
                            apply_font(run, bold=False, size=11, color="000000")

    doc.save(output_path)
    print(f"تم الانتهاء! تم حفظ الملف المنسق في: {output_path}")

if __name__ == "__main__":
    format_document("سياسة_حفظ_الوثائق_مصحح_نهائي.docx", "سياسة_حفظ_الوثائق_منسق_احترافي.docx")
