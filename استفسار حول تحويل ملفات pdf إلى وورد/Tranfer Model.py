import os
import sys
import time
from docx import Document
from mistralai.client import Mistral
from dotenv import load_dotenv

# Fix for printing Arabic text in Windows console
if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')

load_dotenv()
api_key = os.getenv("MISTRAL_API_KEY") or "HLRp9J8Nrf7AYkFGRB6E76JajlGihrhh"
client = Mistral(api_key=api_key)

# نستخدم موديل النصوص السريع والقوي للمراجعة الإملائية
MODEL = "open-mistral-nemo"

CORRECTION_PROMPT = """
أنت مصحح لغوي دقيق.
قم بتصحيح الأخطاء الإملائية والمسافات الخاطئة في النص العربي التالي فقط (مثال: "التمهي د" تصبح "التمهيد").
لا تقم بتغيير المعنى أو إعادة صياغة النص.
إذا كان النص يحتوي على أرقام أو تواريخ أو رموز، حافظ عليها تماماً.
إذا كان النص سليماً أعده كما هو.
أعد النص المصحح فقط بدون أي مقدمات أو علامات اقتباس.
النص:
"""

def correct_text(text, retries=3):
    if not text or len(text.strip()) < 2:
        return text
    
    messages = [{"role": "user", "content": f"{CORRECTION_PROMPT}\n{text}"}]

    for attempt in range(retries):
        try:
            response = client.chat.complete(model=MODEL, messages=messages)
            corrected = response.choices[0].message.content.strip()
            # إزالة علامات الاقتباس إن وجدت
            if corrected.startswith('"') and corrected.endswith('"'):
                corrected = corrected[1:-1]
            return corrected
        except Exception as e:
            err_str = str(e)
            if '429' in err_str or 'rate' in err_str.lower():
                wait = 2 * (attempt + 1)
                print(f"  ⏳ تجاوز حد الطلبات، انتظار {wait} ثانية...")
                time.sleep(wait)
            else:
                print(f"  ❌ خطأ في التصحيح: {e}")
                return text
    return text

def replace_paragraph_text(p, new_text):
    """يستبدل نص الفقرة مع الحفاظ على التنسيق الداخلي بقدر الإمكان"""
    if len(p.runs) == 0:
        p.text = new_text
        return
    # نضع النص المصحح في أول Run ونفرغ الباقي للحفاظ على تنسيق الفقرة (لون/خط)
    p.runs[0].text = new_text
    for i in range(1, len(p.runs)):
        p.runs[i].text = ""

def process_word_document(input_docx, output_docx):
    print(f"بدء التصحيح الإملائي الذكي لملف: {input_docx}")
    doc = Document(input_docx)
    
    # 1. تصحيح الفقرات العادية
    total_paras = len(doc.paragraphs)
    print(f"جاري مراجعة {total_paras} فقرة خارج الجداول...")
    for i, p in enumerate(doc.paragraphs):
        original_text = p.text.strip()
        if original_text:
            corrected = correct_text(original_text)
            if corrected != original_text:
                replace_paragraph_text(p, corrected)

    # 2. تصحيح الجداول
    print("جاري مراجعة النصوص داخل الجداول...")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    original_text = p.text.strip()
                    if original_text:
                        corrected = correct_text(original_text)
                        if corrected != original_text:
                            replace_paragraph_text(p, corrected)
    
    doc.save(output_docx)
    print(f"✅ تم الانتهاء بنجاح! تم حفظ الملف في: {output_docx}")

if __name__ == "__main__":
    INPUT_DOCX = "سيــاســــة-حفظ-الوثائق-و-إتلافها-.docx"
    OUTPUT_DOCX = "سياسة_حفظ_الوثائق_مصحح_نهائي.docx"
    process_word_document(INPUT_DOCX, OUTPUT_DOCX)
