import asyncio
import json
import random
import sys
import os
import pandas as pd
from playwright.async_api import async_playwright
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

sys.stdout.reconfigure(encoding='utf-8')

# 1. قائمة الحسابات الستة المستهدفة بالكامل
TARGET_ACCOUNTS = [
    #"ChinaEmbKSA",      # سفارة الصين في السعودية
    #"CGTNArabic",       # الصين بالعربية
    # "chenweiqing_ksa",  # الدبلوماسي الصيني
    #"USAbilAraby",      # أمريكا بالعربي
    #"CENTCOMArabic",    # القيادة المركزية
    "USEmbassyCairo"    # سفارة أمريكا بالقاهرة
]

# 2. تحديد النطاق الكامل للـ 17 شهراً
START_DATE = "2025-02-01"
END_DATE = "2026-06-30"

# قائمة لتخزين كافة البيانات
all_scraped_data = []

USER_DATA_DIR = "./twitter_profile"

def get_month_ranges(start_date_str, end_date_str):
    import datetime
    import calendar
    start = datetime.datetime.strptime(start_date_str, "%Y-%m-%d")
    end = datetime.datetime.strptime(end_date_str, "%Y-%m-%d")
    
    ranges = []
    current = start
    while current <= end:
        last_day = calendar.monthrange(current.year, current.month)[1]
        month_end = datetime.datetime(current.year, current.month, last_day)
        if month_end > end:
            month_end = end
        ranges.append((current.strftime("%Y-%m-%d"), month_end.strftime("%Y-%m-%d")))
        current = month_end + datetime.timedelta(days=1)
    return ranges

async def check_login_and_wait(page):
    await page.goto("https://x.com/home")
    # إذا تم توجيهنا لصفحة تسجيل الدخول، ننتظر المستخدم ليسجل دخوله
    if "login" in page.url or "i/flow/login" in page.url or page.url == "https://x.com/":
        print("[*] يبدو أنك غير مسجل الدخول.")
        print("[*] الرجاء تسجيل الدخول في النافذة المفتوحة. أمامك 5 دقائق...")
        try:
            import re
            await page.wait_for_url(re.compile(r".*/home.*"), timeout=300000)
            print("[✓] تم تسجيل الدخول وحفظ الجلسة بنجاح!")
        except Exception as e:
            print("[!] انتهى وقت الانتظار أو حدث خطأ. يرجى المحاولة مرة أخرى.")
            return False
    else:
        print("[✓] أنت مسجل الدخول بالفعل. سيتم بدء السحب...")
    return True

def save_to_excel():
    if not all_scraped_data:
        return
        
    df = pd.DataFrame(all_scraped_data)
    output_filename = "Combined_Sample_Data_All_Accounts.xlsx"
    df.to_excel(output_filename, index=False, engine='openpyxl')
    
    # تنسيق ملف الإكسل ليظهر بشكل احترافي
    wb = openpyxl.load_workbook(output_filename)
    ws = wb.active
    ws.sheet_view.rightToLeft = True # تحويل اتجاه الشيت لليمين (عربي)
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F81BD")
    alignment_center = Alignment(horizontal="center", vertical="center")
    alignment_wrap = Alignment(horizontal="right", vertical="center", wrap_text=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = alignment_center
        cell.border = thin_border
        
    column_widths = {
        "A": 22, "B": 20, "C": 45, "D": 15, "E": 20, "F": 15, "G": 15, "H": 15, "I": 80
    }
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width
        
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin_border
            if cell.column == 9: # عامود النص
                cell.alignment = alignment_wrap
            else:
                cell.alignment = alignment_center
                
    wb.save(output_filename)

def load_existing_data():
    output_filename = "Combined_Sample_Data_All_Accounts.xlsx"
    if os.path.exists(output_filename):
        try:
            df = pd.read_excel(output_filename)
            global all_scraped_data
            all_scraped_data = df.to_dict('records')
            print(f"[*] تم العثور على ملف داتا سابق. تم تحميل {len(all_scraped_data)} تغريدة محفوظة مسبقاً لاستكمال السحب عليها.")
        except Exception as e:
            print(f"[!] خطأ في قراءة الملف السابق: {e}")

COMPLETED_MONTHS_FILE = "completed_months.json"
completed_months = []

def load_completed_months():
    global completed_months
    if os.path.exists(COMPLETED_MONTHS_FILE):
        try:
            with open(COMPLETED_MONTHS_FILE, "r", encoding="utf-8") as f:
                completed_months = json.load(f)
        except:
            completed_months = []

def mark_month_completed(account, m_start, m_end):
    key = f"{account}_{m_start}_{m_end}"
    if key not in completed_months:
        completed_months.append(key)
        with open(COMPLETED_MONTHS_FILE, "w", encoding="utf-8") as f:
            json.dump(completed_months, f, ensure_ascii=False)

def is_month_completed(account, m_start, m_end):
    return f"{account}_{m_start}_{m_end}" in completed_months

async def intercept_graphql_and_extract(page, target_account, start_date, end_date):
    async def handle_response(response):
        if "SearchTimeline" in response.url and response.status == 200:
            try:
                raw_data = await response.json()
                tweets = extract_tweets_from_json(raw_data)
                for tweet in tweets:
                    tweet_url = f"https://x.com/{target_account}/status/{tweet.get('rest_id')}"
                    # منع التكرار
                    if not any(d['رابط التغريدة URL'] == tweet_url for d in all_scraped_data):
                        all_scraped_data.append({
                            "تاريخ النشر": tweet.get("created_at"),
                            "الحساب": target_account,
                            "رابط التغريدة URL": tweet_url,
                            "لغة المنشور": tweet.get("lang"),
                            "نوع الوسائط": determine_media_type(tweet),
                            "عدد الإعجابات": tweet.get("favorite_count", 0),
                            "عدد التعليقات": tweet.get("reply_count", 0),
                            "عدد إعادة النشر": tweet.get("retweet_count", 0),
                            "النص المرجعي": tweet.get("full_text", "")
                        })
            except Exception as e:
                pass

    page.on("response", handle_response)

    query = f"from:{target_account} since:{start_date} until:{end_date}"
    search_url = f"https://x.com/search?q={query}&src=typed_query&f=live"
    
    print(f"[*] جاري سحب التغريدات لحساب: {target_account} (من {start_date} إلى {end_date}) ...")
    
    # محاولة تجنب الحظر (Rate Limit) وإعادة التحميل إذا ظهر "Something went wrong"
    for attempt in range(4):
        await page.goto(search_url, wait_until="domcontentloaded", timeout=60000)
        await page.wait_for_timeout(5000)
        
        try:
            page_text = await page.evaluate("document.body.innerText")
            if "Something went wrong" in page_text or "حدث خطأ" in page_text or "Retry" in page_text:
                print(f"[!] تويتر أوقف البحث مؤقتاً (طبيعي جداً بسبب كثرة التغريدات).")
                
                if attempt < 3:
                    print(f"[*] يرجى عدم إيقاف السكربت! السكربت سينتظر دقيقة أو دقيقتين ويعيد المحاولة تلقائياً (المحاولة {attempt+1}/4)...")
                    await page.wait_for_timeout(random.randint(60000, 90000))
                else:
                    # الحد الأقصى للمحاولات السريعة فشل، يجب الانتظار 15 دقيقة (مدة فك الحظر الرسمية لتويتر)
                    print(f"\n[!!!] حسابك وصل للحد الأقصى للاستعلامات (Rate Limit) لدى تويتر.")
                    print(f"[!!!] تويتر يفك هذا الحظر برمجياً بعد 15 دقيقة بالضبط.")
                    print(f"[!!!] السكربت سيدخل الآن في (نوم عميق) لمدة 15 دقيقة ولن يفعل شيئاً.. دعه يعمل ولا توقفه!")
                    
                    # طباعة عد تنازلي كل دقيقة لطمأنة المستخدم
                    for minute in range(15, 0, -1):
                        print(f"[*] متبقي {minute} دقيقة على فك الحظر واستكمال السحب...")
                        await asyncio.sleep(60)
                    
                    print(f"[*] انتهت فترة الحظر! جاري إعادة تحميل الصفحة واستكمال العمل...")
                    await page.goto(search_url, wait_until="domcontentloaded", timeout=60000)
                    await page.wait_for_timeout(5000)
                continue
            else:
                break
        except:
            break
            
    # تحقق أخير بعد كل المحاولات للتأكد من عدم وجود خطأ قبل البدء بالتمرير
    try:
        page_text = await page.evaluate("document.body.innerText")
        if "Something went wrong" in page_text or "حدث خطأ" in page_text or "Retry" in page_text:
            print(f"[!] فشل تخطي الحظر لهذه الفترة ({start_date} إلى {end_date}). سيتم تخطي هذا الشهر لمنع انهيار السكربت بالكامل.")
            return
    except:
        pass

    # حلقة تكرار ذكية للتمرير حتى نهاية البيانات
    previous_height = await page.evaluate("document.body.scrollHeight")
    while True:
        await page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
        # التمرير السريع (سبرنت) لجمع أكبر قدر من الداتا قبل الحظر
        await page.wait_for_timeout(random.randint(3000, 5000))
        
        new_height = await page.evaluate("document.body.scrollHeight")
        # إذا لم يتغير ارتفاع الصفحة، معناه أننا إما وصلنا للنهاية، أو تويتر قام بحظرنا في المنتصف!
        if new_height == previous_height:
            await page.wait_for_timeout(3000)
            final_height = await page.evaluate("document.body.scrollHeight")
            
            # فحص ما إذا كان تويتر قد أظهر رسالة الخطأ (Rate Limit) في منتصف التمرير
            page_text = await page.evaluate("document.body.innerText")
            if "Something went wrong" in page_text or "حدث خطأ" in page_text or "Retry" in page_text:
                print(f"\n[!!!] تويتر أوقف التمرير مؤقتاً (Rate Limit) في منتصف الشهر!")
                print(f"[!!!] السكربت سيدخل في نوم عميق لمدة 15 دقيقة ثم يضغط 'إعادة المحاولة' ويكمل السحب دون أن يفقد التغريدات السابقة...")
                for minute in range(15, 0, -1):
                    print(f"[*] متبقي {minute} دقيقة لاستكمال السحب...")
                    await asyncio.sleep(60)
                
                # البحث عن زر إعادة المحاولة والنقر عليه
                retry_btn = await page.query_selector("button:has-text('Retry')") or await page.query_selector("button:has-text('إعادة المحاولة')")
                if retry_btn:
                    print("[*] انتهى الحظر! جاري النقر على زر إعادة المحاولة...")
                    await retry_btn.click()
                    await page.wait_for_timeout(5000)
                else:
                    print("[*] لم يتم العثور على زر، جاري إعادة تحميل الصفحة...")
                    await page.reload(wait_until="domcontentloaded")
                    await page.wait_for_timeout(5000)
                continue # نعود لدورة التمرير
                
            elif final_height == previous_height:
                print(f"[✓] تم الوصول لنهاية التغريدات المتاحة لحساب {target_account} لهذا الشهر.")
                break
        previous_height = new_height
    
    # إزالة المستمع حتى لا يتداخل مع الحساب التالي
    page.remove_listener("response", handle_response)

def extract_tweets_from_json(json_data):
    tweets = []
    try:
        instructions = json_data['data']['search_by_raw_query']['search_timeline']['timeline']['instructions']
        for instruction in instructions:
            if instruction.get('type') == 'TimelineAddEntries':
                for entry in instruction['entries']:
                    if 'tweet' in entry['entryId']:
                        tweet_core = entry['content']['itemContent']['tweet_results']['result']['legacy']
                        tweet_core['rest_id'] = entry['content']['itemContent']['tweet_results']['result']['rest_id']
                        tweets.append(tweet_core)
    except KeyError:
        pass
    return tweets

def determine_media_type(tweet_legacy):
    if 'extended_entities' in tweet_legacy:
        media = tweet_legacy['extended_entities'].get('media', [])
        if media:
            media_type = media[0].get('type')
            if media_type == 'photo':
                return "صورة / إنفوجرافيك"
            elif media_type == 'video':
                return "فيديو"
            elif media_type == 'animated_gif':
                return "فيديو (GIF)"
    return "نص فقط"

async def main():
    load_existing_data()
    load_completed_months()
    print("[*] تهيئة المتصفح...")
    async with async_playwright() as p:
        # استخدام المتصفح بوضع مرئي لتسجيل الدخول لأول مرة، ومخفي للمرات القادمة
        # هنا سنستخدم وضع مرئي headless=False في كل مرة لكي تستطيع رؤية ما يحدث والتأكد
        context = await p.chromium.launch_persistent_context(
            user_data_dir=USER_DATA_DIR,
            headless=False, 
            channel="chrome",
            viewport={"width": 1280, "height": 720},
            locale="ar-EG"
        )
        
        page = await context.new_page()
        # إضافة سكربت تخطي الحماية
        await page.add_init_script("""
            Object.defineProperty(navigator, 'webdriver', {
                get: () => undefined
            });
        """)
        
        is_logged_in = await check_login_and_wait(page)
        if not is_logged_in:
            await context.close()
            return
            
        print("[*] بدء عملية جمع العينات بطريقة الدفعات الشهرية...")
        
        date_ranges = get_month_ranges(START_DATE, END_DATE)
        
        for account in TARGET_ACCOUNTS:
            print(f"\n==================================================")
            print(f"[*] البدء بحساب: {account}")
            print(f"==================================================")
            for m_start, m_end in date_ranges:
                if is_month_completed(account, m_start, m_end):
                    print(f"[*] تم تخطي ({m_start} إلى {m_end}) لأنه تم سحبه مسبقاً بالكامل.")
                    continue
                    
                await intercept_graphql_and_extract(page, account, m_start, m_end)
                
                # حفظ الداتا فوراً بعد كل شهر لضمان عدم ضياعها
                save_to_excel()
                
                # تسجيل الشهر كـ "مكتمل" لكي لا يعود إليه مرة أخرى لو تم إعادة تشغيل السكربت
                mark_month_completed(account, m_start, m_end)
                
                print(f"[✓] تم حفظ التغريدات بنجاح في الإكسل (الإجمالي حتى الآن: {len(all_scraped_data)}).")
                
                print("[*] انتظار 60 إلى 90 ثانية قبل الشهر القادم لتبريد الحظر (لا توقف السكربت)...")
                await asyncio.sleep(random.randint(60, 90))
            
            # توقف طويل بين كل حساب والذي يليه لتجنب الحظر
            print("[*] انتظار دقيقتين قبل الحساب القادم...")
            await asyncio.sleep(120)
        
        await context.close()
    
    # 4. طباعة النتيجة النهائية
    if all_scraped_data:
        print(f"\n[✓] اكتملت العملية بنجاح! تم سحب وحفظ {len(all_scraped_data)} تغريدة في ملف: Combined_Sample_Data_All_Accounts.xlsx")
    else:
        print("\n[!] لم يتم العثور على بيانات. تأكد من اتصال الإنترنت أو القيود المحتملة.")

if __name__ == "__main__":
    asyncio.run(main())