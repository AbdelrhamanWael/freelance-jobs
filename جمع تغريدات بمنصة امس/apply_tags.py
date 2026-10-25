import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_tag(cell, tag_text):
    """
    يكتب الـ tag في الخلية كـ run واحد متكامل حتى تقدر docxtpl تتعرف عليه.
    بيمسح كل الـ paragraphs الموجودة وبيبني paragraph جديد بـ run واحد.
    """
    tc = cell._tc
    # Remove all existing paragraphs from the cell
    for p in tc.findall(qn('w:p')):
        tc.remove(p)
    
    # Create a new paragraph with a single run
    new_p = OxmlElement('w:p')
    new_r = OxmlElement('w:r')
    new_t = OxmlElement('w:t')
    new_t.text = tag_text
    new_t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    new_r.append(new_t)
    new_p.append(new_r)
    tc.append(new_p)

def apply_tags(doc_path, out_path):
    doc = docx.Document(doc_path)
    
    # Table 4 (البيانات التعريفية)
    t4 = doc.tables[4]
    set_cell_tag(t4.cell(1, 0), "{{Likes}} | {{Comments}} | {{Retweets}}")
    set_cell_tag(t4.cell(1, 1), "{{Media_Type}}")
    set_cell_tag(t4.cell(1, 2), "{{Occasion}}")
    set_cell_tag(t4.cell(1, 3), "{{Language}}")
    set_cell_tag(t4.cell(1, 4), "{{Tweet_URL}}")
    set_cell_tag(t4.cell(1, 6), "{{Account_Name}}")
    set_cell_tag(t4.cell(1, 7), "{{Date}}")

    # التمهيد المنهجي
    for p in doc.paragraphs:
        if "التمهيد المنهجي" in p.text:
            for run in p.runs:
                run.text = ""
            if p.runs:
                p.runs[0].text = "{{Intro_Paragraph}}"
            break

    # Table 5 (التحليل السياقي)
    t5 = doc.tables[5]
    set_cell_tag(t5.cell(2, 0), "{{Qualitative_Analysis_Space}}")
    set_cell_tag(t5.cell(2, 2), "{{Comm_Goal}}")
    set_cell_tag(t5.cell(2, 3), "{{Target_Audience}}")
    set_cell_tag(t5.cell(2, 4), "{{Actors}}")
    set_cell_tag(t5.cell(2, 5), "{{Regional_Context}}")
    set_cell_tag(t5.cell(2, 6), "{{Geo_Scope}}")
    
    # Table 6 (فان دايك)
    t6 = doc.tables[6]
    set_cell_tag(t6.cell(3, 0), "{{Self_Positives_Score}} - {{Self_Positives_Analysis}}")
    set_cell_tag(t6.cell(4, 0), "{{Self_Negatives_Score}} - {{Self_Negatives_Analysis}}")
    set_cell_tag(t6.cell(5, 0), "{{Other_Negatives_Score}} - {{Other_Negatives_Analysis}}")
    set_cell_tag(t6.cell(6, 0), "{{Other_Positives_Score}} - {{Other_Positives_Analysis}}")
    
    # Table 7 (استراتيجيات الدبلوماسية) - بالصف المضبوط مباشرة
    t7 = doc.tables[7]
    t7_tag_map = {
        3:  "{{Official_Info}}",
        4:  "{{Official_Statements}}",
        5:  "{{Intl_Coop}}",
        6:  "{{Diplomatic_Visits}}",
        7:  "{{Stats_Reports}}",
        10: "{{Intl_Sources}}",
        11: "{{Human_Stories}}",
        12: "{{National_Success}}",
        13: "{{State_Superiority}}",
        14: "{{Solidarity}}",
        17: "{{Political_Values}}",
        18: "{{Tech_Progress}}",
        19: "{{Intl_Leadership}}",
        20: "{{Reliable_Partner}}",
        21: "{{Correct_Stereotypes}}",
        22: "{{Intl_Trust}}",
        23: "{{Cultural_Appeal}}",
        26: "{{Hashtags}}",
        27: "{{Mentions}}",
        28: "{{Retweets_Quotes}}",
        29: "{{Call_to_Action}}",
        30: "{{Official_Links}}",
        33: "{{Comms_Channels}}",
        34: "{{Reply_Audience}}",
        35: "{{Live_Sessions}}",
        36: "{{Encourage_Proposals}}",
        37: "{{Polls_Questions}}",
    }
    for row_idx, tag in t7_tag_map.items():
        if row_idx < len(t7.rows):
            set_cell_tag(t7.rows[row_idx].cells[0], tag)

    # Table 8 (التكتيكات الخطابية) - العمود 5 هو الملاحظات
    t8 = doc.tables[8]
    t8_tag_map = {
        3:  "{{Irony}}",
        4:  "{{Celeb_Quotes}}",
        5:  "{{Disapproval}}",
        6:  "{{Politeness}}",
        7:  "{{Show_Power}}",
        8:  "{{Cultural_Symbols}}",
        9:  "{{National_Praise}}",
        10: "{{Endearment}}",
        11: "{{Doubt}}",
        12: "{{Linguistic_Styles}}",
        13: "{{Threats}}",
        14: "{{Victimhood}}",
        15: "{{Similarities}}",
        16: "{{Challenge}}",
        17: "{{Slogans}}",
        18: "{{Citizen_Opinions}}",
        19: "{{Categorization}}",
        20: "{{Positives_Traits}}",
        21: "{{Other_Negatives_Show}}",
        22: "{{Superlatives}}",
        23: "{{Supportive_Power}}",
        26: "{{Examples}}",
        27: "{{Comparison}}",
        28: "{{Proof}}",
        29: "{{Justification}}",
        30: "{{Stats_Numbers}}",
        31: "{{Number_Game}}",
        32: "{{Religious_Frames}}",
    }
    for row_idx, tag in t8_tag_map.items():
        if row_idx < len(t8.rows):
            set_cell_tag(t8.rows[row_idx].cells[5], tag)
                    
    # Table 9 (موارد القوة الناعمة)
    t9 = doc.tables[9]
    t9_tag_map = {
        2: "{{Soft_Pol_Power}}",
        3: "{{Soft_Cult_Power}}",
        4: "{{Soft_Edu_Power}}",
        5: "{{Soft_Econ_Power}}",
        6: "{{Soft_Tech_Power}}",
        7: "{{Soft_Human_Power}}",
        8: "{{Soft_Media_Power}}",
    }
    for row_idx, tag in t9_tag_map.items():
        if row_idx < len(t9.rows):
            set_cell_tag(t9.rows[row_idx].cells[0], tag)

    doc.save(out_path)
    print(f"Template saved to {out_path}")

apply_tags("استمارة-معدلة.docx", "template.docx")
