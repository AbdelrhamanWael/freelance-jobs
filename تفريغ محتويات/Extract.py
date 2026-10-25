"""
Extract.py — Full Design Text Extractor
يستخرج كل النصوص من كل صفحات الـ PDF ويضعها في جدول وورد منظم.
EN في عمود و AR في عمود — بدون أي فلترة أو حذف.
"""

import re
import fitz                              # PyMuPDF
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
import logging
from datetime import datetime
import os

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")

# ══════════════════════════════════════════════════════════════════════════════
#  DESIGN CONSTANTS
# ══════════════════════════════════════════════════════════════════════════════

COLOR_HEADER_BG   = "1F3864"
COLOR_ROW_ODD     = "FFFFFF"
COLOR_ROW_EVEN    = "EEF2F7"
COLOR_BORDER      = "ADB9CA"

FONT_MAIN         = "Arial"
FONT_SIZE_HEADER  = 12
FONT_SIZE_BODY    = 10
FONT_SIZE_TITLE   = 18
FONT_SIZE_SUB     = 10

COL_NO_W  = Cm(1.4)
COL_EN_W  = Cm(13.2)
COL_AR_W  = Cm(13.2)

# ══════════════════════════════════════════════════════════════════════════════
#  TEXT EXTRACTION — بسيط وكامل بدون فلترة
# ══════════════════════════════════════════════════════════════════════════════

ARABIC_RE = re.compile(r'[\u0600-\u06FF]')

def extract_text_from_page(raw_text: str) -> tuple[list[str], list[str]]:
    """
    يفصل النص إلى سطور إنجليزية وعربية.
    لا يحذف أي شيء — كل سطر له نص يُحفظ.
    """
    en_lines = []
    ar_lines = []

    for line in raw_text.splitlines():
        line = line.strip()
        if not line:
            continue
        if ARABIC_RE.search(line):
            ar_lines.append(line)
        else:
            en_lines.append(line)

    return en_lines, ar_lines


# ══════════════════════════════════════════════════════════════════════════════
#  DOCX STYLING HELPERS
# ══════════════════════════════════════════════════════════════════════════════

def set_cell_bg(cell, hex_color: str):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    old = tcPr.find(qn("w:shd"))
    if old is not None:
        tcPr.remove(old)
    tcPr.append(shd)

def set_cell_borders(cell, color: str = COLOR_BORDER):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"),   "single")
        b.set(qn("w:sz"),    "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), color)
        tcBorders.append(b)
    old = tcPr.find(qn("w:tcBorders"))
    if old is not None:
        tcPr.remove(old)
    tcPr.append(tcBorders)

def set_cell_valign(cell, align: str = "center"):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    v    = OxmlElement("w:vAlign")
    v.set(qn("w:val"), align)
    old = tcPr.find(qn("w:vAlign"))
    if old is not None:
        tcPr.remove(old)
    tcPr.append(v)

def set_cell_padding(cell, top=80, bottom=80, left=100, right=100):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for side, val in [("top", top), ("bottom", bottom),
                       ("left", left), ("right", right)]:
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"),    str(val))
        m.set(qn("w:type"), "dxa")
        tcMar.append(m)
    old = tcPr.find(qn("w:tcMar"))
    if old is not None:
        tcPr.remove(old)
    tcPr.append(tcMar)

def set_cell_width(cell, width):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW  = OxmlElement("w:tcW")
    try:
        twips = int(width.twips)
    except Exception:
        twips = int(width)
    tcW.set(qn("w:w"),    str(twips))
    tcW.set(qn("w:type"), "dxa")
    old = tcPr.find(qn("w:tcW"))
    if old is not None:
        tcPr.remove(old)
    tcPr.append(tcW)

def set_run_cs_font(run, font_name: str = FONT_MAIN):
    r   = run._r
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:cs"),    font_name)

def enable_rtl_para(para):
    pPr  = para._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.append(bidi)

def set_row_height(row, height_cm: float):
    trPr     = row._tr.get_or_add_trPr()
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"),   str(int(Cm(height_cm).twips)))
    trHeight.set(qn("w:hRule"), "atLeast")
    old = trPr.find(qn("w:trHeight"))
    if old is not None:
        trPr.remove(old)
    trPr.append(trHeight)

def keep_row_together(row):
    trPr      = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement("w:cantSplit")
    trPr.append(cantSplit)

# ── Cell writers ──────────────────────────────────────────────────────────────

def write_header_cell(cell, text: str, rtl: bool = False):
    cell.text = ""
    set_cell_bg(cell, COLOR_HEADER_BG)
    set_cell_borders(cell, COLOR_HEADER_BG)
    set_cell_valign(cell, "center")
    set_cell_padding(cell, 100, 100, 120, 120)
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if rtl:
        enable_rtl_para(para)
    run = para.add_run(text)
    run.bold = True
    run.font.size      = Pt(FONT_SIZE_HEADER)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    set_run_cs_font(run)

def write_index_cell(cell, text: str, bg: str):
    cell.text = ""
    set_cell_bg(cell, bg)
    set_cell_borders(cell, COLOR_BORDER)
    set_cell_valign(cell, "center")
    set_cell_padding(cell, 80, 80, 60, 60)
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    run.bold           = True
    run.font.size      = Pt(FONT_SIZE_BODY)
    run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    set_run_cs_font(run)

def write_content_cell(cell, lines: list[str], bg: str,
                        align=WD_ALIGN_PARAGRAPH.LEFT, rtl: bool = False):
    """يكتب كل سطر كفقرة منفصلة في الخلية."""
    cell.text = ""
    set_cell_bg(cell, bg)
    set_cell_borders(cell, COLOR_BORDER)
    set_cell_valign(cell, "top")
    set_cell_padding(cell, 90, 90, 130, 130)

    if not lines:
        para = cell.paragraphs[0]
        para.alignment = align
        if rtl:
            enable_rtl_para(para)
        run = para.add_run("—")
        run.font.size      = Pt(FONT_SIZE_BODY)
        run.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
        set_run_cs_font(run)
        return

    first = True
    for line in lines:
        para  = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        para.alignment = align
        if rtl:
            enable_rtl_para(para)
        para.paragraph_format.space_after  = Pt(2)
        para.paragraph_format.space_before = Pt(1)
        run = para.add_run(line)
        run.font.size      = Pt(FONT_SIZE_BODY)
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
        set_run_cs_font(run)

# ══════════════════════════════════════════════════════════════════════════════
#  DOCUMENT SETUP
# ══════════════════════════════════════════════════════════════════════════════

def setup_landscape(doc: Document):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.orientation   = WD_ORIENT.LANDSCAPE
    sec.top_margin    = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin   = Cm(1.5)
    sec.right_margin  = Cm(1.5)

def add_title_block(doc: Document, pdf_path: str, total_rows: int):
    # العنوان الرئيسي
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(0)
    t.paragraph_format.space_after  = Pt(4)
    r = t.add_run("Design Text Extraction — نصوص التصاميم")
    r.bold           = True
    r.font.size      = Pt(FONT_SIZE_TITLE)
    r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    r.font.name      = FONT_MAIN

    # العنوان الفرعي
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.space_before = Pt(0)
    s.paragraph_format.space_after  = Pt(10)
    file_name = os.path.basename(pdf_path)
    date_str  = datetime.now().strftime("%B %d, %Y")
    r2 = s.add_run(
        f"Source: {file_name}    |    Slides extracted: {total_rows}"
        f"    |    Generated: {date_str}"
    )
    r2.font.size      = Pt(FONT_SIZE_SUB)
    r2.font.color.rgb = RGBColor(0x55, 0x6B, 0x8A)
    r2.font.name      = FONT_MAIN

    # خط فاصل
    div = doc.add_paragraph()
    div.paragraph_format.space_after = Pt(8)
    pPr  = div._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot  = OxmlElement("w:bottom")
    bot.set(qn("w:val"),   "single")
    bot.set(qn("w:sz"),    "8")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), COLOR_HEADER_BG)
    pBdr.append(bot)
    pPr.append(pBdr)

def setup_table(table):
    """يضبط الجدول: borders + column widths + header row."""
    table.style = "Table Grid"

    # Table-level borders
    tbl   = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)

    tblBorders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"),   "single")
        b.set(qn("w:sz"),    "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), COLOR_BORDER)
        tblBorders.append(b)
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    tblPr.append(tblBorders)

    # Column grid
    tblGrid = tbl.find(qn("w:tblGrid"))
    if tblGrid is not None:
        tbl.remove(tblGrid)
    tblGrid = OxmlElement("w:tblGrid")
    for w in [COL_NO_W, COL_EN_W, COL_AR_W]:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w.twips)))
        tblGrid.append(gc)
    tbl.insert(0, tblGrid)

    # Header row
    hdr_row = table.rows[0]
    set_row_height(hdr_row, 1.1)
    cells = hdr_row.cells
    write_header_cell(cells[0], "NO")
    write_header_cell(cells[1], "EN TEXT")
    write_header_cell(cells[2], "AR TEXT", rtl=True)

    # تكرار الـ header في كل صفحة
    trPr   = hdr_row._tr.get_or_add_trPr()
    tblHdr = OxmlElement("w:tblHeader")
    trPr.append(tblHdr)


# ══════════════════════════════════════════════════════════════════════════════
#  MAIN FUNCTION
# ══════════════════════════════════════════════════════════════════════════════

def extract_all_text_to_docx(
    pdf_path: str,
    docx_template_path: str,
    output_docx_path: str,
) -> None:

    # ── فتح الـ PDF ──
    try:
        pdf_doc = fitz.open(pdf_path)
    except Exception as e:
        logging.error(f"Cannot open PDF: {e}")
        return

    # ── جمع النصوص من كل الصفحات ──
    pages_data = []
    for page_num in range(len(pdf_doc)):
        page     = pdf_doc.load_page(page_num)
        raw_text = page.get_text("text")
        en_lines, ar_lines = extract_text_from_page(raw_text)

        if not en_lines and not ar_lines:
            logging.info(f"Page {page_num+1}: [empty — skipped]")
            continue

        pages_data.append({
            "page": page_num + 1,
            "en":   en_lines,
            "ar":   ar_lines,
        })
        logging.info(
            f"Page {page_num+1}: {len(en_lines)} EN lines, {len(ar_lines)} AR lines"
        )

    logging.info(f"\nTotal slides with text: {len(pages_data)}")

    # ── فتح القالب ──
    try:
        doc = Document(docx_template_path)
    except Exception as e:
        logging.error(f"Cannot open template: {e}")
        return

    if not doc.tables:
        logging.error("No table in template.")
        return

    # ── إعداد المستند ──
    setup_landscape(doc)
    add_title_block(doc, pdf_path, len(pages_data))

    table = doc.tables[0]
    setup_table(table)

    # ── كتابة الصفوف ──
    for idx, data in enumerate(pages_data, start=1):
        bg = COLOR_ROW_ODD if idx % 2 == 1 else COLOR_ROW_EVEN

        row = table.add_row()
        set_row_height(row, 0.9)
        keep_row_together(row)
        cells = row.cells

        write_index_cell(cells[0], str(idx), bg)
        write_content_cell(
            cells[1], data["en"], bg,
            align=WD_ALIGN_PARAGRAPH.LEFT, rtl=False
        )
        write_content_cell(
            cells[2], data["ar"], bg,
            align=WD_ALIGN_PARAGRAPH.RIGHT, rtl=True
        )

        set_cell_width(cells[0], COL_NO_W)
        set_cell_width(cells[1], COL_EN_W)
        set_cell_width(cells[2], COL_AR_W)

    # ── حفظ ──
    try:
        doc.save(output_docx_path)
        logging.info(f"\nSaved → {output_docx_path}  ({len(pages_data)} slides)")
    except Exception as e:
        logging.error(f"Save failed: {e}")


# ══════════════════════════════════════════════════════════════════════════════
#  ENTRY POINT
# ══════════════════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    PDF_FILE      = "Full Creative Campaign - Downtown Campaign.pdf"
    DOCX_TEMPLATE = "template.docx"
    OUTPUT_FILE   = "Final_Delivery.docx"

    extract_all_text_to_docx(PDF_FILE, DOCX_TEMPLATE, OUTPUT_FILE)