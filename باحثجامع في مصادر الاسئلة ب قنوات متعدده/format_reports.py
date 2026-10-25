import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

def format_excel_report(file_path):
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return
        
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    
    # 1. RTL
    ws.sheet_view.rightToLeft = True
    
    # Fonts
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="0F172A")
    placeholder_font = Font(name="Segoe UI", size=9, italic=True, color="94A3B8")
    
    high_priority_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    high_priority_font = Font(name="Segoe UI", size=10, bold=True, color="991B1B")
    
    normal_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    normal_font = Font(name="Segoe UI", size=10, color="334155")
    
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color="CBD5E1"),
        right=Side(style='thin', color="CBD5E1"),
        top=Side(style='thin', color="CBD5E1"),
        bottom=Side(style='thin', color="CBD5E1")
    )
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)
    
    # Header row height
    ws.row_dimensions[1].height = 32
    
    # Format Headers
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(1, col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border

    screenshot_col_idx = None
    category_col_idx = None
    text_col_idx = None
    link_col_idx = None
    
    for idx, h in enumerate(headers, 1):
        if h == "سكرين شوت":
            screenshot_col_idx = idx
        elif h == "التصنيف":
            category_col_idx = idx
        elif h and "نص الرسالة" in str(h):
            text_col_idx = idx
        elif h and "رابط" in str(h):
            link_col_idx = idx

    # Format Data Rows
    for r in range(2, ws.max_row + 1):
        # Make row comfortably tall for pasting images or PDF icons
        ws.row_dimensions[r].height = 50
        
        category_val = ws.cell(r, category_col_idx).value if category_col_idx else ""
        is_high_priority = (category_val == "أولوية عالية")
        
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.border = thin_border
            
            # Alignments
            if c == text_col_idx:
                cell.alignment = align_right
                cell.font = data_font
            else:
                cell.alignment = align_center
                cell.font = data_font
                
            # Category styling
            if c == category_col_idx:
                if is_high_priority:
                    cell.fill = high_priority_fill
                    cell.font = high_priority_font
                else:
                    cell.fill = normal_fill
                    cell.font = normal_font
            
            # Screenshot column placeholder & hint
            if c == screenshot_col_idx:
                val = str(cell.value or "").strip()
                if val and val != "لا توجد صورة" and val != "[📷 صورة / 📎 ملف PDF]" and os.path.exists(val):
                    abs_file_path = os.path.abspath(val)
                    ext = os.path.splitext(val)[1].lower()
                    if ext in ['.jpg', '.jpeg', '.png', '.webp', '.bmp']:
                        try:
                            from openpyxl.drawing.image import Image as OpenpyxlImage
                            xl_img = OpenpyxlImage(abs_file_path)
                            w, h = xl_img.width, xl_img.height
                            if w > 0 and h > 0:
                                scale = min(130 / w, 50 / h, 1.0)
                                xl_img.width = int(w * scale)
                                xl_img.height = int(h * scale)
                                cell_coord = get_column_letter(c) + str(r)
                                ws.add_image(xl_img, cell_coord)
                            cell.value = ""
                            cell.hyperlink = abs_file_path
                        except Exception:
                            cell.value = f"📷 {os.path.basename(val)}"
                            cell.hyperlink = abs_file_path
                            cell.font = Font(name="Segoe UI", size=10, color="2563EB", underline="single")
                    elif ext == '.pdf':
                        cell.value = f"📄 {os.path.basename(val)}"
                        cell.hyperlink = abs_file_path
                        cell.font = Font(name="Segoe UI", size=10, color="DC2626", underline="single", bold=True)
                    else:
                        cell.value = f"📎 {os.path.basename(val)}"
                        cell.hyperlink = abs_file_path
                        cell.font = Font(name="Segoe UI", size=10, color="2563EB", underline="single")
                elif not cell.value:
                    cell.value = "لا توجد صورة"
                    cell.font = placeholder_font
                cell.alignment = align_center
                
            # Make links clickable
            if c == link_col_idx and cell.value and str(cell.value).startswith("http"):
                cell.hyperlink = cell.value
                cell.font = Font(name="Segoe UI", size=10, color="2563EB", underline="single")

    # Column Widths
    col_widths = {
        "الكلمة المفتاحية": 18,
        "التصنيف": 18,
        "اسم المرسل": 22,
        "تاريخ ووقت الرسالة": 22,
        "سكرين شوت": 28,  # Generous width for image / PDF icon
        "رابط الرسالة": 35,
        "نص الرسالة (للمراجعة)": 55
    }
    
    for idx, h in enumerate(headers, 1):
        col_letter = get_column_letter(idx)
        width = col_widths.get(h, 20)
        ws.column_dimensions[col_letter].width = width

    wb.save(file_path)
    print(f"Successfully formatted: {file_path}")

if __name__ == "__main__":
    reports = [
        "Almufakir/Report_2026-09-08_20-46.xlsx",
        "Elmonsefqudrat/Report_2026-09-08_20-46.xlsx"
    ]
    for r in reports:
        format_excel_report(r)
