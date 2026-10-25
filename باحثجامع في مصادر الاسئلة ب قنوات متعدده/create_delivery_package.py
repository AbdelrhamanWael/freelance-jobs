import os
import sys
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.worksheet.datavalidation import DataValidation

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

delivery_folder = 'تسليم_2026-09-09'
os.makedirs(delivery_folder, exist_ok=True)

# 1. Copy the 3 question bank files into delivery_folder
shutil.copy2('1_بنك_الأسئلة_المجمع_الشامل.xlsx', os.path.join(delivery_folder, '1_بنك_الأسئلة_المجمع_الشامل.xlsx'))
shutil.copy2('2_تقرير_قناة_المفكر_المصور.xlsx', os.path.join(delivery_folder, '2_تقرير_قناة_المفكر_المصور.xlsx'))
shutil.copy2('3_تقرير_قناة_المنصف_المصور.xlsx', os.path.join(delivery_folder, '3_تقرير_قناة_المنصف_المصور.xlsx'))
print('Copied 3 question bank files into delivery folder.')

# 2. Select the top 6 highest-priority students
top_6_students = [
    {
        'sender': '✨👩‍⚕',
        'link': 'https://t.me/Almufakir/25001817',
        'text': 'جاني اليوم في الاختبار (سؤال كمي)',
        'img': 'Almufakir/images/msg_25001808.jpg',
        'date': '2026-09-09 09:24:15'
    },
    {
        'sender': 'Basel Alharthy',
        'link': 'https://t.me/Almufakir/25007161',
        'text': 'جاني اليوم (تأكيد سؤال متكرر)',
        'img': 'Almufakir/images/msg_25007150.jpg',
        'date': '2026-09-09 11:06:50'
    },
    {
        'sender': 'باسل',
        'link': 'https://t.me/Almufakir/25001792',
        'text': 'جاني في الاختبار اول (تأكيد مسألة الحسابات)',
        'img': 'Almufakir/images/msg_25001724.jpg',
        'date': '2026-09-09 09:22:40'
    },
    {
        'sender': 'ل',
        'link': 'https://t.me/Elmonsefqudrat/17263233',
        'text': 'هذا له صيغه ثانيه؟ لان جاني في الاختبار بس ما اذكر ان فيه نسبه في السوال',
        'img': 'Elmonsefqudrat/images/msg_17263175.jpg',
        'date': '2026-09-09 04:36:12'
    },
    {
        'sender': 'L',
        'link': 'https://t.me/Elmonsefqudrat/17260991',
        'text': 'هذا اي بنك؟ جاني باختباري (سؤال مقارنة)',
        'img': 'Elmonsefqudrat/images/msg_17260846.jpg',
        'date': '2026-09-09 03:48:22'
    },
    {
        'sender': 'ً',
        'link': 'https://t.me/Elmonsefqudrat/17276368',
        'text': 'جاني ا ن ج يمكن بس بدون اي تعريف بس قارن بين مساحة المثلث ومساحة الدائرة',
        'img': 'Elmonsefqudrat/images/msg_17276356.jpg',
        'date': '2026-09-09 14:15:30'
    }
]

# 3. Create 4_ملف_التواصل_ومتابعة_الـ_6_طلاب.xlsx
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'متابعة الـ 6 طلاب'
ws.sheet_view.rightToLeft = True

title_font = Font(name='Segoe UI', size=14, bold=True, color='1E3A8A')
note_font = Font(name='Segoe UI', size=10, italic=True, color='475569')
header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
data_font = Font(name='Segoe UI', size=10, color='0F172A')
header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

ws.merge_cells('A1:F1')
ws['A1'] = '📋 ملف التواصل وجمع البيانات - متابعة الـ 6 طلاب الأعلى أولوية'
ws['A1'].font = title_font
ws['A1'].alignment = Alignment(horizontal='right', vertical='center')
ws.row_dimensions[1].height = 28

ws.merge_cells('A2:F2')
ws['A2'] = '📌 ملاحظة: قائمة بالـ 6 طلاب الذين شاركوا تجاربهم وأكدوا ورود الأسئلة في اختبار اليوم، جاهزون لبدء المراسلة وتوثيق الردود.'
ws['A2'].font = note_font
ws['A2'].alignment = Alignment(horizontal='right', vertical='center')
ws.row_dimensions[2].height = 22

headers = [
    'حالة التواصل (تم الرد / في الانتظار)',
    'اسم / يوزر الطالب',
    'رابط حساب الطالب أو رسالته',
    'رد الطالب (الأسئلة والتجربة في الاختبار)',
    'سكرين شوت المحادثة / السؤال',
    'تاريخ ووقت الرسالة'
]
ws.row_dimensions[4].height = 30
for c_idx, h in enumerate(headers, 1):
    c = ws.cell(row=4, column=c_idx, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = thin_border

dv = DataValidation(type='list', formula1='"تم الرد,في الانتظار"', allow_blank=True)
ws.add_data_validation(dv)
dv.add('A5:A20')

status_fill = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
status_font = Font(name='Segoe UI', size=10, bold=True, color='92400E')

for idx, student in enumerate(top_6_students, 5):
    ws.row_dimensions[idx].height = 65
    
    # 1. حالة التواصل
    c1 = ws.cell(row=idx, column=1, value='في الانتظار')
    c1.alignment = Alignment(horizontal='center', vertical='center')
    c1.border = thin_border
    c1.fill = status_fill
    c1.font = status_font
    
    # 2. اسم الطالب
    c2 = ws.cell(row=idx, column=2, value=student['sender'])
    c2.alignment = Alignment(horizontal='center', vertical='center')
    c2.border = thin_border
    c2.font = Font(name='Segoe UI', size=10, bold=True, color='0F172A')
    
    # 3. رابط الرسالة
    c3 = ws.cell(row=idx, column=3, value=student['link'])
    c3.alignment = Alignment(horizontal='center', vertical='center')
    c3.border = thin_border
    c3.hyperlink = student['link']
    c3.font = Font(name='Segoe UI', size=10, color='2563EB', underline='single')
    
    # 4. رد الطالب
    c4 = ws.cell(row=idx, column=4, value=student['text'])
    c4.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
    c4.border = thin_border
    c4.font = data_font
    
    # 5. سكرين شوت
    c5 = ws.cell(row=idx, column=5)
    c5.border = thin_border
    img_path = student['img']
    if img_path and os.path.exists(img_path):
        abs_img = os.path.abspath(img_path)
        c5.value = '📷 عرض المحادثة'
        c5.hyperlink = abs_img
        c5.font = Font(name='Segoe UI', size=9, color='2563EB', underline='single', bold=True)
        c5.alignment = Alignment(horizontal='center', vertical='bottom')
        try:
            xl_img = OpenpyxlImage(abs_img)
            w, h = xl_img.width, xl_img.height
            if w > 0 and h > 0:
                scale = min(110 / w, 44 / h, 1.0)
                xl_img.width = int(w * scale)
                xl_img.height = int(h * scale)
                ws.add_image(xl_img, f'E{idx}')
        except Exception:
            pass
    else:
        c5.value = 'لا توجد صورة'
        c5.font = Font(name='Segoe UI', size=9, italic=True, color='94A3B8')
        c5.alignment = Alignment(horizontal='center', vertical='center')
        
    # 6. التاريخ
    c6 = ws.cell(row=idx, column=6, value=student['date'])
    c6.alignment = Alignment(horizontal='center', vertical='center')
    c6.border = thin_border
    c6.font = data_font

col_widths = {
    1: 25,  # حالة التواصل
    2: 20,  # اسم الطالب
    3: 32,  # رابط الحساب أو الرسالة
    4: 55,  # رد الطالب وتجربته
    5: 25,  # سكرين شوت
    6: 22   # التاريخ
}
for col_idx, width in col_widths.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = width

file4_path = os.path.join(delivery_folder, '4_ملف_التواصل_ومتابعة_الـ_6_طلاب.xlsx')
wb.save(file4_path)

# Also update the root template file
try:
    wb.save('الإكسل الثاني - التواصل وجمع البيانات.xlsx')
    wb.save('التواصل_وجمع_البيانات.xlsx')
except Exception:
    pass

print(f'Successfully created {file4_path} and assembled delivery package!')
