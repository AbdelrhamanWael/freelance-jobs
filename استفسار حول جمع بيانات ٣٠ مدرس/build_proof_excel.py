# -*- coding: utf-8 -*-
"""
توليد وتحديث سجل إثبات المراسلة الموحد (Proof of Outreach):
- يتم التحديث مباشرة في نفس الملف الموحد: qudrat_teachers_outreach_proof.xlsx
- يتضمن الآن 29 معلماً تم إرسال رسائل التعاون لهم بنجاح وموثقين بالسكرين شوت المضمنة
- يوثق رد وموافقة المعلم أ. فهيد الشمري على تقديم العمل المجاني مع منصة يزيد
- يوثق استلام رد فوري من م. كامل (علوم للأبحاث والتصاميم)
- المعلم رقم 30 (أ. عبدالمنعم جوهر) موثق للمكالمات الهاتفية فقط لعدم توفر واتساب على رقمه
- صور السكرين شوت مضمنة مباشرة داخل خلايا الإكسل
- روابط مباشرة لفتح الصور الأصلية بدقة فائقة
- أرقام هواتف تبدأ بـ + وروابط واتساب نشطة
"""

import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils import get_column_letter

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

PROOFS_DIR = "proofs"

OUTREACH_DATA = [
    # 1. أ. محمد الرفاعي
    {
        "id": 1,
        "name": "أ. محمد الرفاعي",
        "platform": "YouTube",
        "specialty": "كمي",
        "phone": "+966530046521",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "09:13 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_mohammad_alrifaii_966530046521_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_mohammad_alrifaii_966530046521.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 2. د. كريم
    {
        "id": 2,
        "name": "د. كريم (qudratDr karim)",
        "platform": "YouTube",
        "specialty": "قدرات (عام)",
        "phone": "+966549172952",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "09:14 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_dr_karim_966549172952_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_dr_karim_966549172952.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 3. أ. إيهاب فهمي
    {
        "id": 3,
        "name": "أ. إيهاب فهمي (GAT4U)",
        "platform": "YouTube",
        "specialty": "قدرات (عام)",
        "phone": "+966569061170",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "09:16 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_ehab_fahmy_966569061170_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_ehab_fahmy_966569061170.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 4. مستر علاء رسلان
    {
        "id": 4,
        "name": "مستر علاء رسلان (MR. RASLAN)",
        "platform": "YouTube",
        "specialty": "قدرات (عام)",
        "phone": "+966565016982",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "09:45 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_mr_raslan_966565016982_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_mr_raslan_966565016982.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 5. أ. سالم زهرة
    {
        "id": 5,
        "name": "أ. سالم زهرة",
        "platform": "YouTube",
        "specialty": "كمي",
        "phone": "+966542372928",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "09:47 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_salem_zahra_966542372928_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_salem_zahra_966542372928.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 6. المدرب تامر الصاوي
    {
        "id": 6,
        "name": "المدرب تامر الصاوي (tamerGAT)",
        "platform": "YouTube",
        "specialty": "قدرات (عام)",
        "phone": "+966553676132",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "09:48 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_tamer_elsawy_966553676132_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_tamer_elsawy_966553676132.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 7. أ. فهيد الشمري (مع الرد والموافقة!)
    {
        "id": 7,
        "name": "أ. فهيد الشمري",
        "platform": "YouTube",
        "specialty": "قدرات (عام)",
        "phone": "+966509950956",
        "status": "🎉 تم الرد والموافقة على التعاون",
        "time_sent": "09:49 PM (الرد 09:51 PM)",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_fahaid_alshammari_reply_966509950956_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_fahaid_alshammari_reply_966509950956.png"),
        "notes": "رد المعلم بالموافقة: (وعليكم السلام .. بإذن الله اتشرف ان اقدم عمل مجاني)"
    },
    # 8. د. محمد الشاعر
    {
        "id": 8,
        "name": "د. محمد الشاعر",
        "platform": "YouTube",
        "specialty": "كمي",
        "phone": "+966567338081",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "09:50 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_dr_mohammad_alshaer_966567338081_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_dr_mohammad_alshaer_966567338081.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 9. أ. إيمان وهدان
    {
        "id": 9,
        "name": "أ. إيمان وهدان (Eman Wahdan)",
        "platform": "YouTube",
        "specialty": "قدرات (عام)",
        "phone": "+966501138137",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "09:54 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_eman_wahdan_966501138137_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_eman_wahdan_966501138137.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 10. أ. روان
    {
        "id": 10,
        "name": "أ. روان (قدرات روان)",
        "platform": "YouTube",
        "specialty": "قدرات (عام)",
        "phone": "+966535266025",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "09:55 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_qudrat_rawan_966535266025_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_qudrat_rawan_966535266025.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 11. أبو عمران
    {
        "id": 11,
        "name": "أبو عمران (قدرات كمي)",
        "platform": "YouTube",
        "specialty": "كمي",
        "phone": "+966536581290",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "09:56 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_abu_emran_966536581290_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_abu_emran_966536581290.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 12. أ. جمال عبدالله
    {
        "id": 12,
        "name": "أ. جمال عبدالله",
        "platform": "YouTube",
        "specialty": "كمي",
        "phone": "+966532860447",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "08:28 AM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_gamal_abdullah_966532860447_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_gamal_abdullah_966532860447.jpg"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 13. أ. أشرف سليمان
    {
        "id": 13,
        "name": "أ. أشرف سليمان (معلم قدرات كمي)",
        "platform": "X (Twitter)",
        "specialty": "كمي",
        "phone": "+966532687910",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "08:29 AM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_ashraf_soliman_966532687910_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_ashraf_soliman_966532687910.jpg"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 14. أ. أمل
    {
        "id": 14,
        "name": "أ. أمل (معلمة قدرات وتحصيلي)",
        "platform": "X (Twitter)",
        "specialty": "كمي + لفظي",
        "phone": "+966581441101",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "08:30 AM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_amal_966581441101_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_amal_966581441101.jpg"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 15. أ. شعبان حمادة (جديد)
    {
        "id": 15,
        "name": "أ. شعبان حمادة",
        "platform": "YouTube",
        "specialty": "لفظي",
        "phone": "+966508121578",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "02:29 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_shaaban_hamada_966508121578_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_shaaban_hamada_966508121578.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 16. أ. رغد الفرشان (جديد)
    {
        "id": 16,
        "name": "أ. رغد الفرشان",
        "platform": "YouTube",
        "specialty": "قدرات (عام)",
        "phone": "+966501490718",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "02:30 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_raghad_alfrshan_966501490718_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_raghad_alfrshan_966501490718.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 17. أ. منى (معلمة قدرات وتحصيلي) (جديد)
    {
        "id": 17,
        "name": "أ. منى (معلمة قدرات وتحصيلي)",
        "platform": "X (Twitter)",
        "specialty": "كمي",
        "phone": "+966560957545",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "02:31 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_mona_966560957545_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_mona_966560957545.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 18. مدرس قدرات ورياضيات جامعي بالرياض (جديد)
    {
        "id": 18,
        "name": "مدرس قدرات ورياضيات جامعي",
        "platform": "YouTube",
        "specialty": "كمي",
        "phone": "+966559704045",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "02:31 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_math_teacher_riyadh_966559704045_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_math_teacher_riyadh_966559704045.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 19. د. سحر أحمد
    {
        "id": 19,
        "name": "د. سحر أحمد (Dr. Sahar Ahmed)",
        "platform": "X (Twitter)",
        "specialty": "كمي + لفظي",
        "phone": "+966599064204",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "02:37 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_sahar_ahmed_966599064204_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_sahar_ahmed_966599064204.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 20. أ. عمر
    {
        "id": 20,
        "name": "أ. عمر (أستاذ اللفظي)",
        "platform": "X (Twitter)",
        "specialty": "لفظي",
        "phone": "+966557347832",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "02:38 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_omar_lafzi_966557347832_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_omar_lafzi_966557347832.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 21. م. كامل (علوم للأبحاث والتصاميم)
    {
        "id": 21,
        "name": "م. كامل (معلم قدرات وإحصاء - علوم)",
        "platform": "X (Twitter)",
        "specialty": "كمي",
        "phone": "+966530605399",
        "status": "✅ تم الإرسال بنجاح (استلام رد آلي)",
        "time_sent": "02:39 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_kamel_3lom_966530605399_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_kamel_3lom_966530605399.png"),
        "notes": "تم إرسال الدعوة والملف + استلام رد فوري: (شكرا لتواصلك مع علوم للأبحاث والتصاميم.. كيف يمكننا خدمتك)"
    },
    # 22. د. ماث
    {
        "id": 22,
        "name": "د. ماث (Dr Math - قدرات كمي)",
        "platform": "X (Twitter)",
        "specialty": "كمي",
        "phone": "+966565046450",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "02:40 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_dr_math_966565046450_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_dr_math_966565046450.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 23. أ. رنا
    {
        "id": 23,
        "name": "أ. رنا (Rana - قدرات)",
        "platform": "X (Twitter)",
        "specialty": "قدرات (عام)",
        "phone": "+966545385974",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "02:40 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_rana_qudrat_966545385974_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_rana_qudrat_966545385974.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 24. أ. ساره
    {
        "id": 24,
        "name": "أ. ساره (تدريب قدرات ورخصة)",
        "platform": "X (Twitter)",
        "specialty": "كمي",
        "phone": "+966581019236",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "02:42 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_sarah_qudrat_966581019236_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_sarah_qudrat_966581019236.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 25. مدرس قدرات وتحصيلي بالرياض
    {
        "id": 25,
        "name": "مدرس قدرات وتحصيلي بالرياض",
        "platform": "X (Twitter)",
        "specialty": "قدرات (عام)",
        "phone": "+966548434100",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "02:42 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_qdratqiyas4u_riyadh_966548434100_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_qdratqiyas4u_riyadh_966548434100.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 26. معلم قدرات كمي (قناة قدرات كمي)
    {
        "id": 26,
        "name": "معلم قدرات كمي (قناة قدرات كمي)",
        "platform": "YouTube",
        "specialty": "كمي",
        "phone": "+966505330678",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "02:56 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_qudrat_kammi_channel_966505330678_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_qudrat_kammi_channel_966505330678.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 27. أ. محمد رضا (قاف بلس)
    {
        "id": 27,
        "name": "أ. محمد رضا (قاف بلس)",
        "platform": "YouTube",
        "specialty": "لفظي",
        "phone": "+966504910099",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "02:56 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_mohammad_reda_qafplus_966504910099_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_mohammad_reda_qafplus_966504910099.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 28. dr trainer (مدرس قدرات بالرياض)
    {
        "id": 28,
        "name": "dr trainer (مدرس قدرات بالرياض)",
        "platform": "YouTube",
        "specialty": "قدرات (عام)",
        "phone": "+966541872325",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "02:57 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_dr_trainer_gat_966541872325_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_dr_trainer_gat_966541872325.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 29. معلم قدرات كمي بالخرج
    {
        "id": 29,
        "name": "معلم قدرات كمي بالخرج",
        "platform": "X (Twitter)",
        "specialty": "كمي",
        "phone": "+966542795845",
        "status": "✅ تم الإرسال بنجاح",
        "time_sent": "02:57 PM",
        "thumb_image": os.path.join(PROOFS_DIR, "proof_teacher_alkharj_966542795845_thumb.png"),
        "full_image": os.path.join(PROOFS_DIR, "proof_teacher_alkharj_966542795845.png"),
        "notes": "تم إرسال نص دعوة التعاون لمنصة يزيد + ملف PDF الجزء الأول (419 سؤال)"
    },
    # 30. أ. عبدالمنعم جوهر
    {
        "id": 30,
        "name": "أ. عبدالمنعم جوهر",
        "platform": "YouTube",
        "specialty": "لفظي",
        "phone": "+966549257154",
        "status": "📞 اتصال هاتفي فقط (لا يتوفر واتساب)",
        "time_sent": "-",
        "thumb_image": None,
        "full_image": None,
        "notes": "الرقم متاح للمكالمات الهاتفية المباشرة (غير مفعل عليه واتساب) - جاهز للتواصل عبر الاتصال"
    }
]

def build_proof_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "سجل إثبات المراسلة (Proof)"
    ws.views.sheetView[0].rightToLeft = True

    headers = [
        "م", 
        "اسم المعلم", 
        "المنصة", 
        "التخصص", 
        "رقم الواتساب (+)", 
        "حالة الإرسال (Proof)", 
        "وقت الإرسال", 
        "صورة الإثبات (سكرين شوت)", 
        "رابط الصورة بدقة عالية",
        "تفاصيل الرسالة والملف المرسل"
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Segoe UI", size=10)
    success_font = Font(name="Segoe UI", size=10, bold=True, color="047857")
    approval_font = Font(name="Segoe UI", size=10, bold=True, color="1E40AF") # Deep blue for approval
    phone_call_font = Font(name="Segoe UI", size=10, bold=True, color="B45309") # Amber for phone call
    pending_font = Font(name="Segoe UI", size=10, color="6B7280")
    phone_font = Font(name="Segoe UI", size=10, bold=True, color="059669", underline="single")
    link_font = Font(name="Segoe UI", size=10, color="1D4ED8", underline="single")
    
    success_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    approval_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid") # soft blue for approval
    phone_call_fill = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid") # light warm amber
    pending_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    ws.row_dimensions[1].height = 32
    for c_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for r_idx, row_data in enumerate(OUTREACH_DATA, start=2):
        is_approval = "🎉" in row_data["status"]
        is_sent = "✅" in row_data["status"]
        is_call = "📞" in row_data["status"]
        
        if is_approval:
            row_fill = approval_fill
        elif is_sent:
            row_fill = success_fill
        elif is_call:
            row_fill = phone_call_fill
        else:
            row_fill = pending_fill
        
        if row_data["thumb_image"] and os.path.exists(row_data["thumb_image"]):
            ws.row_dimensions[r_idx].height = 125
        else:
            ws.row_dimensions[r_idx].height = 28
            
        full_img_text = "🔎 عرض صورة الإثبات كاملة" if row_data["full_image"] else "-"
        status_thumb_text = "" if row_data["thumb_image"] else ("مكالمة هاتفية" if is_call else "قيد المراسلة")
        
        ws.append([
            row_data["id"],
            row_data["name"],
            row_data["platform"],
            row_data["specialty"],
            row_data["phone"],
            row_data["status"],
            row_data["time_sent"],
            status_thumb_text,
            full_img_text,
            row_data["notes"]
        ])

        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = body_font
            cell.border = thin_border
            cell.fill = row_fill
            
            if c_idx in [1, 3, 4, 5, 6, 7, 8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
                
            if c_idx == 5:
                clean_phone = row_data["phone"].lstrip("+")
                cell.number_format = "@"
                cell.value = row_data["phone"]
                cell.hyperlink = f"https://wa.me/{clean_phone}"
                cell.font = phone_font

            if c_idx == 6:
                if is_approval:
                    cell.font = approval_font
                elif is_sent:
                    cell.font = success_font
                elif is_call:
                    cell.font = phone_call_font
                else:
                    cell.font = pending_font

            if c_idx == 9 and row_data["full_image"]:
                rel_path = os.path.abspath(row_data["full_image"])
                cell.hyperlink = f"file:///{rel_path.replace(os.sep, '/')}"
                cell.font = link_font

        # تضمين الصورة المصغرة في العمود H
        if row_data["thumb_image"] and os.path.exists(row_data["thumb_image"]):
            try:
                img = ExcelImage(row_data["thumb_image"])
                if "amal" in row_data["thumb_image"] or "ashraf" in row_data["thumb_image"] or "gamal" in row_data["thumb_image"]:
                    img.width = 90
                    img.height = 155
                else:
                    img.width = 175
                    img.height = 145
                cell_coord = f"H{r_idx}"
                ws.add_image(img, cell_coord)
            except Exception as e:
                print(f"Error embedding image for row {r_idx}: {e}")

    col_widths = {
        "A": 8,   # م
        "B": 28,  # اسم المعلم
        "C": 15,  # المنصة
        "D": 16,  # التخصص
        "E": 18,  # رقم الواتساب
        "F": 32,  # حالة الإرسال
        "G": 18,  # وقت الإرسال
        "H": 28,  # صورة الإثبات
        "I": 26,  # رابط الصورة
        "J": 45   # تفاصيل الرسالة
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # حفظ في نفس الملف الموحد مباشرة
    output_file = "qudrat_teachers_outreach_proof.xlsx"
    wb.save(output_file)
    print(f"[✓] تم حفظ وتحديث الملف الموحد بنجاح: {output_file}")

    # تصدير نسخة CSV المرافقة
    import csv
    with open("qudrat_teachers_outreach_proof.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for d in OUTREACH_DATA:
            w.writerow([
                d["id"],
                d["name"],
                d["platform"],
                d["specialty"],
                d["phone"],
                d["status"],
                d["time_sent"],
                d["thumb_image"] if d["thumb_image"] else "-",
                d["full_image"] if d["full_image"] else "-",
                d["notes"]
            ])
    print("[✓] تم حفظ: qudrat_teachers_outreach_proof.csv")

if __name__ == "__main__":
    build_proof_excel()
